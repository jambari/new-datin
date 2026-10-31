import datetime
from django.test import TestCase, Client
from django.urls import reverse
from django.contrib.auth.models import User
from .models import MagneticObservation, Precursor, MagnetDataAvailability


def make_observation(**kwargs):
    defaults = dict(
        observation_date=datetime.date(2024, 3, 1),
        observer='Jambari',
        session='Pagi',
    )
    defaults.update(kwargs)
    return MagneticObservation.objects.create(**defaults)


class MagneticObservationModelTest(TestCase):
    def test_create_and_str(self):
        obs = make_observation()
        self.assertIn('2024-03-01', str(obs))
        self.assertIn('Jambari', str(obs))

    def test_is_bartington_false_without_readings(self):
        obs = make_observation()
        self.assertFalse(obs.is_bartington)

    def test_is_bartington_true_with_deg_readings(self):
        obs = make_observation(deklinasi_readings={
            'WU': {'deg': 10, 'min': 30, 'sec': 0},
        })
        self.assertTrue(obs.is_bartington)

    def test_derived_components_calculated_on_save(self):
        obs = make_observation(
            declination=1.0,
            inclination=-30.0,
            total_intensity=45000.0,
        )
        obs.refresh_from_db()
        self.assertIsNotNone(obs.horizontal_intensity)
        self.assertIsNotNone(obs.vertical_intensity)
        self.assertIsNotNone(obs.north_component)
        self.assertIsNotNone(obs.east_component)

    def test_bartington_full_calculation(self):
        obs = make_observation(
            deklinasi_readings={
                'WU': {'deg': 181, 'min': 0, 'sec': 0},
                'ED': {'deg': 181, 'min': 0, 'sec': 0},
                'WD': {'deg': 181, 'min': 0, 'sec': 0},
                'EU': {'deg': 181, 'min': 0, 'sec': 0},
            },
            inklinasi_readings={
                'NU': {'deg': 150, 'min': 0, 'sec': 0, 'ftotal': 45000},
                'SD': {'deg': 300, 'min': 0, 'sec': 0, 'ftotal': 45000},
                'ND': {'deg': 210, 'min': 0, 'sec': 0, 'ftotal': 45000},
                'SU': {'deg': 60, 'min': 0, 'sec': 0, 'ftotal': 45000},
            },
        )
        obs.refresh_from_db()
        self.assertIsNotNone(obs.declination)
        self.assertIsNotNone(obs.inclination)
        self.assertEqual(float(obs.total_intensity), 45000.0)


class PrecursorModelTest(TestCase):
    def test_create_and_str(self):
        p = Precursor.objects.create(
            anomaly_timestamp=datetime.datetime(2024, 1, 1, tzinfo=datetime.timezone.utc),
            predicted_start_date=datetime.date(2024, 1, 10),
            predicted_end_date=datetime.date(2024, 1, 20),
            predicted_magnitude=6.0,
            location_description='Papua Tengah',
        )
        self.assertIn('Papua Tengah', str(p))

    def test_default_not_validated(self):
        p = Precursor.objects.create(
            anomaly_timestamp=datetime.datetime(2024, 1, 1, tzinfo=datetime.timezone.utc),
            predicted_start_date=datetime.date(2024, 1, 10),
            predicted_end_date=datetime.date(2024, 1, 20),
            predicted_magnitude=6.0,
            location_description='Test',
        )
        self.assertFalse(p.is_validated)


class MagnetDataAvailabilityModelTest(TestCase):
    def test_create_and_str(self):
        obj = MagnetDataAvailability.objects.create(
            station='JYP_P',
            date=datetime.date(2024, 3, 1),
            percentage=98.5,
        )
        self.assertIn('JYP_P', str(obj))

    def test_unique_together_constraint(self):
        MagnetDataAvailability.objects.create(
            station='JYP_P', date=datetime.date(2024, 3, 1), percentage=98.5
        )
        from django.db import IntegrityError
        with self.assertRaises(IntegrityError):
            MagnetDataAvailability.objects.create(
                station='JYP_P', date=datetime.date(2024, 3, 1), percentage=50.0
            )


class MagnetViewsTest(TestCase):
    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user('testuser', password='pass123')
        self.client.login(username='testuser', password='pass123')

    def test_observation_form_get(self):
        resp = self.client.get(reverse('observation_form'))
        self.assertEqual(resp.status_code, 200)

    def test_observation_bartington_form_get(self):
        resp = self.client.get(reverse('observation_bartington'))
        self.assertEqual(resp.status_code, 200)

    def test_observation_list_get(self):
        resp = self.client.get(reverse('observation_list'))
        self.assertEqual(resp.status_code, 200)

    def test_precursor_list_get(self):
        resp = self.client.get(reverse('precursor_list'))
        self.assertEqual(resp.status_code, 200)

    def test_magnet_availability_query_get(self):
        resp = self.client.get(reverse('magnet_availability_query'))
        self.assertEqual(resp.status_code, 200)

    def test_observation_detail_404_on_invalid(self):
        resp = self.client.get(reverse('observation_detail', args=[9999]))
        self.assertEqual(resp.status_code, 404)

    def test_precursor_detail_404_on_invalid(self):
        resp = self.client.get(reverse('precursor_detail', args=[9999]))
        self.assertEqual(resp.status_code, 404)

    def test_unauthenticated_redirects(self):
        self.client.logout()
        resp = self.client.get(reverse('observation_form'))
        self.assertIn(resp.status_code, [302, 403])

    def test_update_magnet_availability_post(self):
        payload = {'station': 'JYP_P', 'date': '2024-03-01', 'value': 99.0}
        resp = self.client.post(
            reverse('update_magnet_availability'),
            data=json.dumps(payload),
            content_type='application/json',
        )
        self.assertIn(resp.status_code, [200, 201])


import json


# ── Indeks magnetbumi K & A (scraping harian) ────────────────────────────────

import hashlib
import shutil
import tempfile
from unittest import mock

from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import override_settings

from .models import FmiIndicesImage

# PNG minimal (header valid) — cukup untuk pengecekan magic bytes di command.
TINY_PNG = (
    b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00'
    b'\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00'
    b'\x00\x00IEND\xaeB`\x82'
)

_CMD = 'magnet.management.commands.fetch_fmi_indices'


class FmiIndicesImageModelTest(TestCase):
    def setUp(self):
        # Jangan tulis ke MEDIA_ROOT sungguhan saat test.
        self.media_root = tempfile.mkdtemp(prefix='fmi-test-media-')
        self.addCleanup(shutil.rmtree, self.media_root, True)
        ctx = override_settings(MEDIA_ROOT=self.media_root)
        ctx.enable()
        self.addCleanup(ctx.disable)

    def test_create_and_str(self):
        obj = FmiIndicesImage.objects.create(
            jenis=FmiIndicesImage.K,
            tanggal=datetime.date(2026, 10, 3),
            image=SimpleUploadedFile('k.png', TINY_PNG, content_type='image/png'),
        )
        self.assertIn('K', str(obj))
        self.assertIn('2026-10-03', str(obj))
        self.assertEqual(obj.jenis_label, 'K Indices')

    def test_unique_per_jenis_and_tanggal(self):
        from django.db import IntegrityError

        FmiIndicesImage.objects.create(
            jenis=FmiIndicesImage.K, tanggal=datetime.date(2026, 10, 3),
            image=SimpleUploadedFile('a.png', TINY_PNG),
        )
        with self.assertRaises(IntegrityError):
            FmiIndicesImage.objects.create(
                jenis=FmiIndicesImage.K, tanggal=datetime.date(2026, 10, 3),
                image=SimpleUploadedFile('b.png', TINY_PNG),
            )

    def test_k_and_a_can_share_the_same_date(self):
        for jenis in (FmiIndicesImage.K, FmiIndicesImage.A):
            FmiIndicesImage.objects.create(
                jenis=jenis, tanggal=datetime.date(2026, 10, 3),
                image=SimpleUploadedFile(f'{jenis}.png', TINY_PNG),
            )
        self.assertEqual(FmiIndicesImage.objects.filter(tanggal=datetime.date(2026, 10, 3)).count(), 2)

    def test_filename_property(self):
        obj = FmiIndicesImage(jenis='A', tanggal=datetime.date(2026, 10, 3))
        self.assertEqual(obj.filename, 'FMI-A-Indices-JYP-20261003.png')


class FmiIndicesFetchCommandTest(TestCase):
    def setUp(self):
        self.media_root = tempfile.mkdtemp(prefix='fmi-test-media-')
        self.addCleanup(shutil.rmtree, self.media_root, True)
        ctx = override_settings(MEDIA_ROOT=self.media_root)
        ctx.enable()
        self.addCleanup(ctx.disable)

    @staticmethod
    def _response(content=TINY_PNG, status=200, headers=None):
        resp = mock.Mock()
        resp.status_code = status
        resp.content = content
        resp.headers = headers if headers is not None else {'ETag': '"etag-1"', 'Content-Type': 'image/png'}
        return resp

    def test_fetch_stores_one_record_per_jenis(self):
        with mock.patch(f'{_CMD}.requests.get', return_value=self._response()):
            call_command('fetch_fmi_indices', '--date', '2026-10-03', '--no-telegram')

        self.assertEqual(FmiIndicesImage.objects.count(), 2)
        expected_sha = hashlib.sha256(TINY_PNG).hexdigest()
        for rec in FmiIndicesImage.objects.all():
            self.assertEqual(rec.tanggal, datetime.date(2026, 10, 3))
            self.assertEqual(rec.size_bytes, len(TINY_PNG))
            self.assertEqual(rec.sha256, expected_sha)
            self.assertTrue(rec.image.storage.exists(rec.image.name))
            self.assertIsNone(rec.telegram_sent_at)  # --no-telegram

    def test_rerun_same_day_is_idempotent(self):
        with mock.patch(f'{_CMD}.requests.get', return_value=self._response()):
            call_command('fetch_fmi_indices', '--date', '2026-10-03', '--no-telegram')
            call_command('fetch_fmi_indices', '--date', '2026-10-03', '--no-telegram')

        self.assertEqual(FmiIndicesImage.objects.count(), 2)
        names = sorted(FmiIndicesImage.objects.values_list('image', flat=True))
        self.assertEqual(len(set(names)), 2)  # tidak ada file bersuffix

    def test_single_jenis_flag(self):
        with mock.patch(f'{_CMD}.requests.get', return_value=self._response()):
            call_command('fetch_fmi_indices', '--date', '2026-10-03', '--jenis', 'K', '--no-telegram')
        self.assertEqual(FmiIndicesImage.objects.count(), 1)
        self.assertEqual(FmiIndicesImage.objects.first().jenis, 'K')

    def test_non_png_response_is_rejected(self):
        bad = self._response(content=b'<html>not an image</html>', headers={'Content-Type': 'text/html'})
        with mock.patch(f'{_CMD}.requests.get', return_value=bad), \
             mock.patch(f'{_CMD}.time.sleep', return_value=None):
            with self.assertRaises(CommandError):
                call_command('fetch_fmi_indices', '--date', '2026-10-03', '--no-telegram')
        self.assertEqual(FmiIndicesImage.objects.count(), 0)

    def test_telegram_failure_does_not_block_storage(self):
        """Token kosong => pengiriman dilewati, tetapi citra tetap tersimpan."""
        with override_settings(TELEGRAM_BOT_TOKEN='', TELEGRAM_CHAT_ID=''), \
             mock.patch(f'{_CMD}.requests.get', return_value=self._response()):
            call_command('fetch_fmi_indices', '--date', '2026-10-03')
        self.assertEqual(FmiIndicesImage.objects.count(), 2)
        self.assertIsNone(FmiIndicesImage.objects.first().telegram_sent_at)


class FmiIndicesViewTest(TestCase):
    def setUp(self):
        # Jangan tulis ke MEDIA_ROOT sungguhan saat test.
        self.media_root = tempfile.mkdtemp(prefix='fmi-test-media-')
        self.addCleanup(shutil.rmtree, self.media_root, True)
        ctx = override_settings(MEDIA_ROOT=self.media_root)
        ctx.enable()
        self.addCleanup(ctx.disable)

    @staticmethod
    def _make(jenis, tanggal=datetime.date(2026, 10, 3)):
        """Buat record dengan nama file deterministik seperti hasil command."""
        obj = FmiIndicesImage(jenis=jenis, tanggal=tanggal)
        obj.image.save(obj.filename, ContentFile(TINY_PNG), save=False)
        obj.size_bytes = len(TINY_PNG)
        obj.sha256 = hashlib.sha256(TINY_PNG).hexdigest()
        obj.save()
        return obj

    def test_list_get(self):
        resp = self.client.get(reverse('fmi_indices_list'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Indeks K')

    def test_list_shows_stored_images(self):
        self._make(FmiIndicesImage.K)
        self._make(FmiIndicesImage.A)
        resp = self.client.get(reverse('fmi_indices_list'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'FMI-K-Indices-JYP-20261003.png')
        self.assertContains(resp, 'FMI-A-Indices-JYP-20261003.png')

    def test_list_empty_state(self):
        resp = self.client.get(reverse('fmi_indices_list'))
        self.assertContains(resp, 'Belum ada citra indeks')

    def test_filter_by_jenis(self):
        self._make(FmiIndicesImage.K)
        resp = self.client.get(reverse('fmi_indices_list'), {'jenis': 'A'})
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, 'FMI-K-Indices-JYP-20261003.png')

    def test_post_fetch_does_not_error_when_broker_unavailable(self):
        with mock.patch('magnet.tasks.fetch_fmi_indices_task.delay', side_effect=RuntimeError('broker down')):
            resp = self.client.post(reverse('fmi_indices_list'), {'action': 'fetch'}, follow=True)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Gagal menjadwalkan job')

    def test_tabs_and_kindex_table(self):
        from .models import KIndexDaily
        self._make(FmiIndicesImage.K)
        KIndexDaily.objects.create(tanggal=datetime.date(2026, 10, 3), k_values=[1, 2, 3, 4, 5, 2, 1, 0],
                                   sumber=KIndexDaily.SUMBER_FMI)
        KIndexDaily.objects.create(tanggal=datetime.date(2026, 10, 4), k_values=[3, None, None, None, None, None, None, None],
                                   sumber=KIndexDaily.SUMBER_FMI)
        KIndexDaily.objects.create(tanggal=datetime.date(2026, 9, 1), k_values=[2] * 8)
        resp = self.client.get(reverse('fmi_indices_list'))
        self.assertContains(resp, 'Tabel Data Indeks K')
        self.assertContains(resp, 'FMI-K-Indices-JYP-20261003.png')
        self.assertNotContains(resp, 'Indeks K per 3 jam (UTC)</th>')

        resp = self.client.get(reverse('fmi_indices_list'), {'tab': 'tabel', 'start_date': '2026-10-01'})
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, 'FMI-K-Indices-JYP-20261003.png')
        rows = resp.context['table_rows']
        self.assertEqual([r['obj'].tanggal.day for r in rows], [4, 3])      # terbaru dulu, filter tanggal
        self.assertEqual(rows[1]['sum_k'], 18)
        self.assertEqual([lvl for _k, lvl in rows[1]['cells']][3:5], ['active', 'storm'])
        self.assertContains(resp, '(1/8)')                                  # indeks A sementara
        self.assertContains(resp, 'Citra FMI')


class FmiIndicesPublicPageTest(TestCase):
    """Halaman publik indeks K & A + /magnetbumi/ + submenu navbar landing."""

    def setUp(self):
        self.media_root = tempfile.mkdtemp(prefix='fmi-test-media-')
        self.addCleanup(shutil.rmtree, self.media_root, True)
        ctx = override_settings(MEDIA_ROOT=self.media_root)
        ctx.enable()
        self.addCleanup(ctx.disable)

    @staticmethod
    def _make(jenis, tanggal=datetime.date(2026, 10, 3)):
        obj = FmiIndicesImage(jenis=jenis, tanggal=tanggal)
        obj.image.save(obj.filename, ContentFile(TINY_PNG), save=False)
        obj.size_bytes = len(TINY_PNG)
        obj.sha256 = hashlib.sha256(TINY_PNG).hexdigest()
        obj.save()
        return obj

    # ── halaman khusus indeks K & A ──────────────────────────────────────────

    def test_indices_page_shows_latest_pair(self):
        self._make(FmiIndicesImage.K)
        self._make(FmiIndicesImage.A)

        resp = self.client.get(reverse('public_magnetbumi_indices'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'FMI-K-Indices-JYP-20261003.png')
        self.assertContains(resp, 'FMI-A-Indices-JYP-20261003.png')
        self.assertContains(resp, 'K Indices')
        self.assertContains(resp, 'A Indices')

    def test_indices_page_shows_only_latest_date(self):
        self._make(FmiIndicesImage.K, tanggal=datetime.date(2026, 10, 3))
        self._make(FmiIndicesImage.A, tanggal=datetime.date(2026, 10, 3))
        self._make(FmiIndicesImage.K, tanggal=datetime.date(2026, 10, 2))
        self._make(FmiIndicesImage.A, tanggal=datetime.date(2026, 10, 2))

        resp = self.client.get(reverse('public_magnetbumi_indices'))
        self.assertContains(resp, 'FMI-K-Indices-JYP-20261003.png')
        self.assertNotContains(resp, 'FMI-K-Indices-JYP-20261002.png')

    def test_indices_page_empty_state_without_data(self):
        resp = self.client.get(reverse('public_magnetbumi_indices'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Citra indeks K &amp; A belum tersedia')

    def test_indices_page_partial_pair_does_not_break(self):
        """Hanya citra K tersimpan — halaman tetap tampil."""
        self._make(FmiIndicesImage.K)
        resp = self.client.get(reverse('public_magnetbumi_indices'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Citra A belum tersedia')

    def test_indices_page_has_no_magnetbumi_charts(self):
        """Halaman indeks tidak boleh memuat chart 7 komponen."""
        self._make(FmiIndicesImage.K)
        resp = self.client.get(reverse('public_magnetbumi_indices'))
        self.assertNotContains(resp, 'chart-grid')
        self.assertNotContains(resp, 'chart-D')

    def test_indices_page_has_single_shared_navbar(self):
        """Semua halaman publik memakai satu navbar bersama
        (partials/_landing_header.html); header milik halaman sendiri dulu
        menumpuk di atasnya dan terlihat DOBEL."""
        resp = self.client.get(reverse('public_magnetbumi_indices'))
        self.assertContains(resp, '<header id="site-header"', count=1)
        self.assertContains(resp, 'id="pnNav"', count=1)

    # ── /magnetbumi/ tetap 7 komponen, tanpa citra indeks ────────────────────

    def test_magnetbumi_page_has_seven_component_charts(self):
        resp = self.client.get(reverse('public_magnetbumi'))
        self.assertEqual(resp.status_code, 200)
        for canvas in ['chart-D', 'chart-I', 'chart-F', 'chart-H', 'chart-Z', 'chart-X', 'chart-Y']:
            self.assertContains(resp, canvas)

    def test_magnetbumi_page_has_no_index_images(self):
        """Citra K & A tidak lagi ditampilkan di /magnetbumi/ (terpisah)."""
        self._make(FmiIndicesImage.K)
        self._make(FmiIndicesImage.A)

        resp = self.client.get(reverse('public_magnetbumi'))
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, 'id="indeks-k-a"')
        self.assertNotContains(resp, 'FMI-K-Indices-JYP-20261003.png')
        self.assertNotContains(resp, 'FMI-A-Indices-JYP-20261003.png')

    # ── navbar ───────────────────────────────────────────────────────────────

    def test_landing_navbar_links_to_separate_pages(self):
        resp = self.client.get(reverse('landing'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Absolut Magnetik')
        self.assertContains(resp, 'K dan A Indeks')
        # masing-masing menuju halaman sendiri, bukan anchor di halaman yang sama
        self.assertContains(resp, 'href="/magnetbumi/"')
        self.assertContains(resp, 'href="/magnetbumi/indeks-k-a/"')
        self.assertNotContains(resp, '/magnetbumi/#indeks-k-a')


# ── Notifikasi Telegram ──────────────────────────────────────────────────────
from unittest import mock

from django.test import override_settings


@override_settings(TELEGRAM_BOT_TOKEN='tok-123', TELEGRAM_CHAT_ID='-100123')
class BartingtonTelegramNotifyTest(TestCase):
    """Observasi absolut Bartington baru dikabarkan ke Telegram."""

    READINGS = {'WU': {'deg': 10, 'min': 30, 'sec': 0}}

    def _create(self, **kwargs):
        with mock.patch('theme.telegram.requests.post') as post:
            post.return_value = mock.Mock(status_code=200, text='')
            with self.captureOnCommitCallbacks(execute=True):
                obs = make_observation(**kwargs)
        return obs, post

    def test_bartington_record_is_sent(self):
        obs, post = self._create(
            observer='Jambari', session='Pagi',
            deklinasi_readings=self.READINGS,
            declination=1.0, inclination=-30.0, total_intensity=45000.0)
        self.assertEqual(post.call_count, 1)

        payload = post.call_args.kwargs['data']
        self.assertEqual(payload['chat_id'], '-100123')
        text = payload['text']
        self.assertIn('BARTINGTON', text)
        self.assertIn('01-03-2024', text)
        self.assertIn('Jambari', text)
        self.assertIn('Pagi', text)
        self.assertIn('45000', text)          # F (nT)

    def test_mingeo_record_is_not_sent(self):
        """Form MinGeo menulis model yang sama tapi satuannya grad."""
        obs, post = self._create()            # tanpa pembacaan 'deg'
        self.assertFalse(obs.is_bartington)
        self.assertEqual(post.call_count, 0)

    def test_update_does_not_notify(self):
        obs, _ = self._create(deklinasi_readings=self.READINGS)
        with mock.patch('theme.telegram.requests.post') as post:
            with self.captureOnCommitCallbacks(execute=True):
                obs.session = 'Sore'
                obs.save()
        self.assertEqual(post.call_count, 0)

    def test_observer_name_is_escaped(self):
        _, post = self._create(observer='A <b>X</b>', deklinasi_readings=self.READINGS)
        self.assertIn('A &lt;b&gt;X&lt;/b&gt;', post.call_args.kwargs['data']['text'])

    @override_settings(TELEGRAM_BOT_TOKEN='', TELEGRAM_CHAT_ID='')
    def test_no_message_when_telegram_not_configured(self):
        _, post = self._create(deklinasi_readings=self.READINGS)
        self.assertEqual(post.call_count, 0)

    def test_formatting_error_never_breaks_the_save(self):
        with mock.patch('magnet.signals.format_bartington_message',
                        side_effect=ValueError('boom')):
            with mock.patch('theme.telegram.requests.post') as post:
                with self.captureOnCommitCallbacks(execute=True):
                    obs = make_observation(deklinasi_readings=self.READINGS)
        self.assertTrue(MagneticObservation.objects.filter(pk=obs.pk).exists())
        self.assertEqual(post.call_count, 0)


# ── Indeks K: urai citra FMI & batas sumber Excel/FMI ────────────────────────

import io
import os

from openpyxl import Workbook

from .fmi_kindex import FmiLayoutError, parse_k_image, update_kindex_from_image
from .kindex_import import import_kindex_workbook
from .models import KIndexDaily

_TESTDATA = os.path.join(os.path.dirname(__file__), 'testdata')
UTC = datetime.timezone.utc


def _testimage(name):
    with open(os.path.join(_TESTDATA, name), 'rb') as f:
        return f.read()


def _by_day(slots):
    days = {}
    for t, k in slots:
        days.setdefault(t.date(), [None] * 8)[t.hour // 3] = k
    return days


class FmiKIndexParseTest(TestCase):
    """Nilai acuan dibaca manual dari citra BMKG asli di magnet/testdata."""

    def test_parse_storm_image(self):
        slots = parse_k_image(io.BytesIO(_testimage('FMI-K-Indices-JYP-20261017.png')), datetime.date(2026, 10, 17))
        self.assertEqual(len(slots), 48)
        self.assertEqual(slots[0][0], datetime.datetime(2026, 10, 11, 3))
        days = _by_day(slots)
        self.assertEqual(days[datetime.date(2026, 10, 11)], [None, 9, 9, 9, 9, 7, 7, 7])  # K 9 tertutup logo
        self.assertEqual(days[datetime.date(2026, 10, 12)], [4, 4, 4, 1, 2, 2, 3, 3])
        self.assertEqual(days[datetime.date(2026, 10, 13)], [None, 4, 3, 3, 3, 2, 3, 3])  # slot kosong
        self.assertEqual(days[datetime.date(2026, 10, 17)][0], 3)

    def test_parse_quiet_image_with_missing_day(self):
        days = _by_day(parse_k_image(io.BytesIO(_testimage('FMI-K-Indices-JYP-20261003.png')), datetime.date(2026, 10, 3)))
        self.assertEqual(days[datetime.date(2026, 9, 27)], [None, 3, 3, 1, 2, 1, 3, 3])
        self.assertEqual(days[datetime.date(2026, 9, 28)], [None] * 8)
        self.assertEqual(days[datetime.date(2026, 9, 30)], [4, 4, 3, 1, 2, 1, 3, 3])

    def test_unexpected_layout_raises(self):
        buf = io.BytesIO()
        from PIL import Image
        Image.new('RGB', (800, 400), 'white').save(buf, 'PNG')
        buf.seek(0)
        with self.assertRaises(FmiLayoutError):
            parse_k_image(buf, datetime.date(2026, 10, 17))


@override_settings(KINDEX_FMI_START='2026-10-03')
class FmiKIndexUpdateTest(TestCase):
    def setUp(self):
        self.media_root = tempfile.mkdtemp(prefix='fmi-test-media-')
        self.addCleanup(shutil.rmtree, self.media_root, True)
        ctx = override_settings(MEDIA_ROOT=self.media_root)
        ctx.enable()
        self.addCleanup(ctx.disable)

    def _image(self, name, downloaded):
        obj = FmiIndicesImage.objects.create(jenis='K', tanggal=downloaded.date(),
                                             image=SimpleUploadedFile(name, _testimage(name)))
        FmiIndicesImage.objects.filter(pk=obj.pk).update(updated_at=downloaded)
        obj.refresh_from_db()
        return obj

    def test_only_days_from_cutoff_and_excel_days_untouched(self):
        KIndexDaily.objects.create(tanggal=datetime.date(2026, 9, 30), k_values=[5, 5, 3, 1, 1, 0, 0, 4])
        img = self._image('FMI-K-Indices-JYP-20261003.png', datetime.datetime(2026, 10, 3, 7, 0, tzinfo=UTC))
        days = update_kindex_from_image(img)
        self.assertEqual(days, [datetime.date(2026, 10, 3)])
        sep30 = KIndexDaily.objects.get(tanggal=datetime.date(2026, 9, 30))
        self.assertEqual(sep30.k_values, [5, 5, 3, 1, 1, 0, 0, 4])
        self.assertEqual(sep30.sumber, KIndexDaily.SUMBER_EXCEL)
        oct3 = KIndexDaily.objects.get(tanggal=datetime.date(2026, 10, 3))
        self.assertEqual(oct3.sumber, KIndexDaily.SUMBER_FMI)
        self.assertEqual(oct3.k_values[0], 4)  # batang terakhir citra (kuning)
        self.assertEqual(oct3.k_values[1:], [None] * 7)
        self.assertFalse(oct3.is_complete)
        self.assertEqual(oct3.a_index, 27)  # rata-rata dari slot yang ada (K4 → 27)

    def test_newer_image_overrides_but_blank_slots_do_not_erase(self):
        KIndexDaily.objects.create(tanggal=datetime.date(2026, 10, 13), sumber='fmi',
                                   k_values=[5, 1, 1, 1, 1, 1, 1, 1])
        img = self._image('FMI-K-Indices-JYP-20261017.png', datetime.datetime(2026, 10, 17, 7, 0, tzinfo=UTC))
        update_kindex_from_image(img)
        oct13 = KIndexDaily.objects.get(tanggal=datetime.date(2026, 10, 13))
        self.assertEqual(oct13.k_values, [5, 4, 3, 3, 3, 2, 3, 3])  # slot 0 kosong di citra → nilai lama tetap
        oct11 = KIndexDaily.objects.get(tanggal=datetime.date(2026, 10, 11))
        self.assertEqual(oct11.k_max, 9)
        self.assertEqual(oct11.criteria, 'Badai Besar')

    def test_fetch_command_parses_k_image(self):
        resp = mock.Mock(status_code=200, content=_testimage('FMI-K-Indices-JYP-20261017.png'),
                         headers={'ETag': '"e"'})
        out = io.StringIO()
        with mock.patch(f'{_CMD}.requests.get', return_value=resp), \
                mock.patch('django.utils.timezone.now',
                           return_value=datetime.datetime(2026, 10, 17, 7, 0, tzinfo=UTC)):
            call_command('fetch_fmi_indices', '--jenis', 'K', '--no-telegram', stdout=out)
        self.assertIn('nilai indeks K diurai', out.getvalue())
        self.assertEqual(KIndexDaily.objects.get(tanggal=datetime.date(2026, 10, 12)).k_values, [4, 4, 4, 1, 2, 2, 3, 3])

    def test_fetch_command_survives_unparseable_image(self):
        resp = mock.Mock(status_code=200, content=TINY_PNG, headers={'ETag': '"e"'})
        err = io.StringIO()
        with mock.patch(f'{_CMD}.requests.get', return_value=resp):
            call_command('fetch_fmi_indices', '--jenis', 'K', '--no-telegram', stderr=err)
        self.assertIn('gagal mengurai', err.getvalue())
        self.assertEqual(FmiIndicesImage.objects.count(), 1)  # citra tetap tersimpan

    def test_backfill_command_in_download_order(self):
        self._image('FMI-K-Indices-JYP-20261003.png', datetime.datetime(2026, 10, 3, 7, 0, tzinfo=UTC))
        self._image('FMI-K-Indices-JYP-20261017.png', datetime.datetime(2026, 10, 17, 7, 0, tzinfo=UTC))
        out = io.StringIO()
        call_command('parse_fmi_kindex', stdout=out)
        self.assertIn('Selesai', out.getvalue())
        self.assertTrue(KIndexDaily.objects.filter(tanggal=datetime.date(2026, 10, 3)).exists())
        self.assertFalse(KIndexDaily.objects.filter(tanggal__lt=datetime.date(2026, 10, 3)).exists())

    def test_excel_import_skips_days_from_cutoff(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Oktober'
        ws['K2'] = 2026
        for day in range(1, 32):
            ws.cell(4 + day, 1, day)
            for j in range(8):
                ws.cell(4 + day, 2 + j, 2)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        summary, warnings = import_kindex_workbook(buf, 2026, 'k.xlsx')
        self.assertEqual(summary['2026-10']['data'], 2)  # hanya 1–2 Oktober
        self.assertTrue(any('dilewati' in w for w in warnings))
        self.assertEqual(KIndexDaily.objects.get(tanggal=datetime.date(2026, 10, 2)).sumber, 'excel')
        self.assertFalse(KIndexDaily.objects.filter(tanggal=datetime.date(2026, 10, 3)).exists())


@override_settings(KINDEX_FMI_START='2026-10-03')
class FmiStaleSourceTest(TestCase):
    """Sumber BMKG tidak diperbarui: jangan buat tanggal baru, jangan geser nilai K."""

    LAST_MODIFIED = 'Sat, 17 Oct 2026 05:20:24 GMT'

    def setUp(self):
        self.media_root = tempfile.mkdtemp(prefix='fmi-test-media-')
        self.addCleanup(shutil.rmtree, self.media_root, True)
        ctx = override_settings(MEDIA_ROOT=self.media_root)
        ctx.enable()
        self.addCleanup(ctx.disable)

    def _fetch(self, day, content=None, etag='"e1"', *args):
        resp = mock.Mock(status_code=200, content=content or _testimage('FMI-K-Indices-JYP-20261017.png'),
                         headers={'ETag': etag, 'Last-Modified': self.LAST_MODIFIED})
        out = io.StringIO()
        with mock.patch(f'{_CMD}.requests.get', return_value=resp), \
                mock.patch(f'{_CMD}._send_photo', return_value=(True, 'ok')) as photo, \
                mock.patch(f'{_CMD}.send_telegram_message', return_value=True) as warn, \
                mock.patch('django.utils.timezone.now',
                           return_value=datetime.datetime(2026, 10, day, 7, 0, tzinfo=UTC)):
            call_command('fetch_fmi_indices', '--jenis', 'K', *args, stdout=out)
        return out.getvalue(), photo, warn

    def test_k_values_follow_last_modified_not_download_day(self):
        self._fetch(19)
        img = FmiIndicesImage.objects.get()
        self.assertEqual(img.source_modified_at, datetime.datetime(2026, 10, 17, 5, 20, 24, tzinfo=UTC))
        self.assertEqual(KIndexDaily.objects.get(tanggal=datetime.date(2026, 10, 12)).k_values,
                         [4, 4, 4, 1, 2, 2, 3, 3])
        self.assertFalse(KIndexDaily.objects.filter(tanggal__gt=datetime.date(2026, 10, 17)).exists())

    def test_unchanged_source_is_not_stored_again_and_warns_once(self):
        _, photo, _ = self._fetch(17)
        self.assertEqual(photo.call_count, 1)
        k_before = list(KIndexDaily.objects.order_by('tanggal').values_list('tanggal', 'k_values'))

        # Byte berbeda (kompresi ulang di sumber) tapi ETag sama → tetap dianggap citra lama.
        out, photo, warn = self._fetch(18, content=_testimage('FMI-K-Indices-JYP-20261017.png') + b'x')
        self.assertIn('sumber belum diperbarui', out)
        self.assertEqual(FmiIndicesImage.objects.count(), 1)
        photo.assert_not_called()
        warn.assert_called_once()
        self.assertIn('17 October 2026', warn.call_args[0][0])
        self.assertEqual(list(KIndexDaily.objects.order_by('tanggal').values_list('tanggal', 'k_values')), k_before)

        _, photo, warn = self._fetch(19)          # hari kedua macet: tanpa peringatan lagi
        photo.assert_not_called()
        warn.assert_not_called()

    def test_new_etag_is_stored(self):
        self._fetch(17)
        self._fetch(18, None, '"e2"')
        self.assertEqual(FmiIndicesImage.objects.count(), 2)

    def test_manual_button_does_not_send_telegram(self):
        with mock.patch('magnet.tasks.fetch_fmi_indices_task.delay') as delay:
            self.client.post(reverse('fmi_indices_list'), {'action': 'fetch'})
        self.assertIs(delay.call_args.kwargs['telegram'], False)
        with mock.patch('magnet.tasks.call_command') as cmd:
            from .tasks import fetch_fmi_indices_task
            fetch_fmi_indices_task(telegram=False)
        self.assertTrue(cmd.call_args.kwargs['no_telegram'])

    def _stored(self, day, etag):
        name = 'FMI-K-Indices-JYP-20261017.png'
        obj = FmiIndicesImage.objects.create(jenis='K', tanggal=datetime.date(2026, 10, day), etag=etag,
                                             image=SimpleUploadedFile(name, _testimage(name)))
        FmiIndicesImage.objects.filter(pk=obj.pk).update(
            updated_at=datetime.datetime(2026, 10, day, 7, 0, tzinfo=UTC))

    def test_rebuild_skips_repeated_etag_and_resets(self):
        self._stored(17, '"e1"')
        self._stored(18, '"e1"')                  # duplikat: sumber macet
        KIndexDaily.objects.create(tanggal=datetime.date(2026, 10, 18), sumber='fmi', k_values=[1] * 8)

        out = io.StringIO()
        call_command('parse_fmi_kindex', '--reset', '--dry-run', stdout=out)
        self.assertIn('dry-run', out.getvalue())
        self.assertEqual(KIndexDaily.objects.get(tanggal=datetime.date(2026, 10, 18)).k_values, [1] * 8)

        out = io.StringIO()
        call_command('parse_fmi_kindex', '--reset', stdout=out)
        self.assertIn('dilewati', out.getvalue())
        self.assertFalse(KIndexDaily.objects.filter(tanggal=datetime.date(2026, 10, 18)).exists())
        self.assertEqual(KIndexDaily.objects.get(tanggal=datetime.date(2026, 10, 12)).k_values,
                         [4, 4, 4, 1, 2, 2, 3, 3])


class FmiKIndexGoldK5Test(TestCase):
    def test_k5_gold_bars_are_read(self):
        """K 5 digambar emas (255,204,0); dulu terbaca kosong."""
        days = _by_day(parse_k_image(io.BytesIO(_testimage('FMI-K-Indices-JYP-20261024.png')),
                                     datetime.date(2026, 10, 24)))
        self.assertEqual(days[datetime.date(2026, 10, 19)], [4, 4, 1, 1, 1, 1, 2, None])
        self.assertEqual(days[datetime.date(2026, 10, 20)], [None, None, None, 3, 3, 3, 3, 5])
        self.assertEqual(days[datetime.date(2026, 10, 21)], [5, 4, 2, 2, 2, 3, 3, 3])


# ── Indeks K hitung lokal (metode FMI) dari data mentah LEMI JYP ─────────────────────

def _write_lemi_day(folder, day, x_func, y_func, step=2):
    """Tulis berkas mentah LEMI sintetis (sampel tiap `step` detik) untuk satu hari UTC."""
    import numpy as np
    t = np.arange(0, 86400, step)
    x, y = x_func(day, t), y_func(day, t)
    lines = [f'{day:%Y %m %d} {s // 3600:02d} {s % 3600 // 60:02d} {s % 60:02d}  {a:.2f} {b:.2f} -14790.00  29.5  28.8 P N FL  93'
             for s, a, b in zip(t, x, y)]
    path = os.path.join(folder, f'{day:%Y %m %d} 00 00 00.txt')
    with open(path, 'w') as fh:
        fh.write('\n'.join(lines) + '\n')
    return path


def _quiet_x(day, t):
    import numpy as np
    return 35780 + 20 * np.sin(2 * np.pi * (t / 86400 - 0.1))


def _quiet_y(day, t):
    import numpy as np
    return -1820 + 8 * np.cos(2 * np.pi * t / 86400)


@override_settings(KINDEX_FMI_START='2026-10-03', KINDEX_JYP_L9=240.0)
class KIndexLocalFmiTest(TestCase):
    def setUp(self):
        self.raw = tempfile.mkdtemp(prefix='lemi-raw-')
        self.addCleanup(shutil.rmtree, self.raw, True)
        ctx = override_settings(LEMI_RAW_DIR=self.raw)
        ctx.enable()
        self.addCleanup(ctx.disable)

    def _days(self, disturbed=None, step_day=None):
        import numpy as np

        def x_func(day, t):
            x = _quiet_x(day, t)
            if day == disturbed:                         # badai 12–15 UTC: osilasi 150 nT
                block = (t >= 12 * 3600) & (t < 15 * 3600)
                x = x + np.where(block, 75 * np.sin(2 * np.pi * t / 1200), 0)
            return x

        def y_func(day, t):
            y = _quiet_y(day, t)
            if day == step_day:                          # perawatan sensor 09:00: Y turun 3000 nT
                y = y + np.where(t < 9 * 3600, 3000, 0)
            elif step_day and day < step_day:
                y = y + 3000
            return y

        for d in (9, 10, 11):
            _write_lemi_day(self.raw, datetime.date(2026, 10, d), x_func, y_func)

    def test_quiet_day_is_low_and_storm_block_is_high(self):
        from .kindex_fmi import compute_day
        self._days(disturbed=datetime.date(2026, 10, 10))
        result = compute_day(datetime.date(2026, 10, 10))
        ks = result['k_values']
        self.assertEqual(len(ks), 8)
        self.assertGreaterEqual(ks[4], 7)                 # 150 nT dengan L9 240 → K7
        self.assertTrue(all(k <= 2 for i, k in enumerate(ks) if i != 4), ks)
        self.assertFalse(result['provisional'])

    def test_baseline_step_is_removed(self):
        from .kindex_fmi import compute_day
        self._days(step_day=datetime.date(2026, 10, 10))
        result = compute_day(datetime.date(2026, 10, 10))
        self.assertEqual(result['steps'], 1)
        self.assertTrue(all(k is not None and k <= 2 for k in result['k_values']), result['k_values'])

    def test_store_day_marks_last_day_provisional_and_keeps_bmkg_parser_out(self):
        from .fmi_kindex import update_kindex_from_image
        from .kindex_fmi import store_day
        self._days()
        _, obj = store_day(datetime.date(2026, 10, 11))  # 12 Okt sudah lewat tapi datanya tidak ada
        self.assertEqual(obj.sumber, KIndexDaily.SUMBER_LOKAL)
        self.assertIn('hari berikutnya tidak ada', obj.keterangan)
        before = list(obj.k_values)

        name = 'FMI-K-Indices-JYP-20261017.png'
        media = tempfile.mkdtemp(prefix='fmi-test-media-')
        self.addCleanup(shutil.rmtree, media, True)
        with override_settings(MEDIA_ROOT=media):
            img = FmiIndicesImage.objects.create(jenis='K', tanggal=datetime.date(2026, 10, 17),
                                                 image=SimpleUploadedFile(name, _testimage(name)))
            FmiIndicesImage.objects.filter(pk=img.pk).update(
                updated_at=datetime.datetime(2026, 10, 17, 7, 0, tzinfo=UTC))
            img.refresh_from_db()
            update_kindex_from_image(img)                # citra BMKG mencakup 11–17 Okt
        self.assertEqual(KIndexDaily.objects.get(tanggal=datetime.date(2026, 10, 11)).k_values, before)
        self.assertEqual(KIndexDaily.objects.get(tanggal=datetime.date(2026, 10, 12)).sumber, 'fmi')

    def test_excel_days_are_not_touched(self):
        from .kindex_fmi import store_day
        self._days()
        KIndexDaily.objects.create(tanggal=datetime.date(2026, 10, 2), k_values=[5] * 8)
        _, obj = store_day(datetime.date(2026, 10, 2))
        self.assertIsNone(obj)
        self.assertEqual(KIndexDaily.objects.get(tanggal=datetime.date(2026, 10, 2)).sumber, 'excel')

    def test_command_and_png_and_public_page(self):
        self._days()
        out = io.StringIO()
        call_command('compute_kindex_jyp', '--start', '2026-10-10', '--end', '2026-10-11', stdout=out)
        self.assertIn('2 hari disimpan', out.getvalue())
        resp = self.client.get(reverse('kindex_jyp_png', args=['k']))
        self.assertEqual(resp['Content-Type'], 'image/png')
        self.assertTrue(resp.content.startswith(b'\x89PNG'))
        resp = self.client.get(reverse('kindex_jyp_png', args=['a']))
        self.assertTrue(resp.content.startswith(b'\x89PNG'))
        page = self.client.get(reverse('public_magnetbumi_indices'))
        self.assertContains(page, 'Hitung lokal dari data LEMI-018 JYP')
        self.assertContains(page, reverse('kindex_jyp_png', args=['k']))


@override_settings(LEMI_MONITOR_TOKEN='rahasia')
class LemiRawApiTest(TestCase):
    def setUp(self):
        self.raw = tempfile.mkdtemp(prefix='lemi-raw-')
        self.addCleanup(shutil.rmtree, self.raw, True)
        ctx = override_settings(LEMI_RAW_DIR=self.raw)
        ctx.enable()
        self.addCleanup(ctx.disable)
        self.url = reverse('lemi_raw_api')
        self.auth = {'HTTP_AUTHORIZATION': 'Bearer rahasia'}
        today = datetime.datetime.now(UTC).date()
        self.name = f'{today:%Y %m %d} 00 00 00.txt'

    def _post(self, content, name=None):
        upload = SimpleUploadedFile(name or self.name, content, content_type='text/plain')
        with mock.patch('magnet.tasks.compute_kindex_jyp_task.delay') as delay:
            resp = self.client.post(self.url, {'file': upload}, **self.auth)
        return resp, delay

    def test_requires_token(self):
        self.assertEqual(self.client.get(self.url).status_code, 401)
        self.assertEqual(self.client.get(self.url, HTTP_AUTHORIZATION='Bearer salah').status_code, 401)

    def test_upload_grow_and_manifest(self):
        resp, delay = self._post(b'baris1\n')
        self.assertTrue(resp.json()['stored'])
        delay.assert_called_once()
        resp, _ = self._post(b'baris1\nbaris2\n')          # berkas hari berjalan bertambah
        self.assertTrue(resp.json()['stored'])
        resp, delay = self._post(b'baris1\n')              # versi lebih kecil tidak menimpa
        self.assertFalse(resp.json()['stored'])
        delay.assert_not_called()
        manifest = self.client.get(self.url, **self.auth).json()['files']
        self.assertEqual(manifest, {self.name: len(b'baris1\nbaris2\n')})

    def test_rejects_bad_names(self):
        for bad in ('../../etc/passwd', 'data.txt', '2026 10 30 00 00 00.exe'):
            resp, _ = self._post(b'x', name=bad)
            self.assertEqual(resp.status_code, 400, bad)
        self.assertEqual(os.listdir(self.raw), [])


# ── Data terpotong: level melintasi celah, kelengkapan, kartu CF, peringatan ─────────

class KIndexCutDataTest(TestCase):
    @staticmethod
    def _quiet_minutes():
        import numpy as np
        t = np.arange(4320) * 60
        return (35780 + 20 * np.sin(2 * np.pi * (t / 86400 - 0.1)),
                -1820 + 8 * np.cos(2 * np.pi * t / 86400))

    def test_offset_across_outage_gives_no_false_storm(self):
        import numpy as np
        from .kindex_fmi import k_fmi_day, level_across_gaps
        x, y = self._quiet_minutes()
        base = k_fmi_day(x, y, 240)[0]
        y2 = y.copy()
        y2[1440 + 360:1440 + 720] = np.nan
        y2[1440 + 720:] += 300
        _, steps = level_across_gaps(y2)
        self.assertEqual(len(steps), 1)
        self.assertAlmostEqual(steps[0][2], -300, delta=5)
        ks = k_fmi_day(x, y2, 240)[0]
        self.assertEqual(ks[2:4], [None, None])
        for a, b in zip(base, ks):
            if b is not None:
                self.assertLessEqual(abs(a - b), 1)

    def test_natural_change_across_gap_is_left_alone(self):
        import numpy as np
        from .kindex_fmi import level_across_gaps
        x, _ = self._quiet_minutes()
        x = x.copy()
        x[1440 + 360:1440 + 720] = np.nan
        x[1440 + 720:] -= 100                       # terbesar yang tercatat di arsip
        fixed, steps = level_across_gaps(x)
        self.assertEqual(steps, [])
        self.assertAlmostEqual(np.nansum(fixed), np.nansum(x))

    def test_coverage_and_labels(self):
        import numpy as np
        from .kindex_fmi import block_coverage
        x, y = self._quiet_minutes()
        x = x.copy()
        x[1440:1440 + 90] = np.nan
        self.assertEqual(block_coverage(x, y), [50, 100, 100, 100, 100, 100, 100, 100])

        obj = KIndexDaily(tanggal=datetime.date(2026, 10, 28), k_values=[5, None, None, None, 1, 2, 2, 3],
                          k_coverage=[50, 0, 0, 0, 100, 100, 100, 100], sumber='lokal', a_index=10)
        self.assertEqual(obj.criteria_label, 'Data kurang')
        self.assertEqual(obj.criteria, 'Tenang')      # laporan bulanan tetap memakai ini
        self.assertEqual(obj.partial_blocks, [0])
        full = KIndexDaily(tanggal=datetime.date(2026, 10, 29), k_values=[1] * 8, a_index=3)
        self.assertEqual(full.criteria_label, 'Tenang')

    def test_past_day_without_next_day_is_not_provisional_forever(self):
        from .kindex_fmi import compute_day
        raw = tempfile.mkdtemp(prefix='lemi-raw-')
        self.addCleanup(shutil.rmtree, raw, True)
        _write_lemi_day(raw, datetime.date(2026, 10, 9), _quiet_x, _quiet_y)
        _write_lemi_day(raw, datetime.date(2026, 10, 10), _quiet_x, _quiet_y)
        result = compute_day(datetime.date(2026, 10, 10), raw_dir=raw, l9=240.0)
        self.assertFalse(result['provisional'])
        self.assertTrue(result['next_day_absent'])
        with mock.patch('magnet.kindex_fmi.datetime') as dt_mock:
            dt_mock.date, dt_mock.timedelta, dt_mock.timezone = datetime.date, datetime.timedelta, datetime.timezone
            dt_mock.datetime.now.return_value = datetime.datetime(2026, 10, 11, 1, tzinfo=UTC)
            result = compute_day(datetime.date(2026, 10, 10), raw_dir=raw, l9=240.0)
        self.assertTrue(result['provisional'])


class LemiFlashTest(TestCase):
    @staticmethod
    def _flash_file(folder, start, sectors):
        import struct
        out = bytearray()
        to_bcd = lambda v: (v // 10) << 4 | (v % 10)
        for s in range(sectors):
            t = start + datetime.timedelta(seconds=30 * s)
            head = b'LEMI018\x81' + bytes(to_bcd(v) for v in (t.year - 2000, t.month, t.day, t.hour, t.minute, t.second))
            out += head.ljust(32, b'\0')
            for i in range(30):
                out += struct.pack('<iiihh', 35780000 + i, -1820000, -14790000, 2950, 2880)
        path = os.path.join(folder, 'L018_001.B81')
        with open(path, 'wb') as fh:
            fh.write(out)
        return path

    def test_fills_only_missing_seconds_and_is_idempotent(self):
        from .kindex_fmi import raw_files_for, read_lemi_raw
        from .lemi_flash import fill_day, read_flash
        raw = tempfile.mkdtemp(prefix='lemi-raw-')
        self.addCleanup(shutil.rmtree, raw, True)
        day = datetime.date(2026, 10, 28)
        with open(os.path.join(raw, '2026 10 28 00 00 00.txt'), 'w') as fh:
            for s in range(30):
                fh.write(f'2026 10 28 00 00 {s:02d}  35000.00  -1000.00 -14000.00  29.00  28.00 P N FL  93\n')
        records = read_flash(self._flash_file(raw, datetime.datetime(2026, 10, 28), 2))[day]
        self.assertEqual(len(records), 60)
        self.assertAlmostEqual(records[31][0], 35780.001)
        self.assertEqual(fill_day(raw, day, records, dry_run=True), 30)
        self.assertEqual(fill_day(raw, day, records), 30)
        self.assertEqual(fill_day(raw, day, records), 0)
        files = raw_files_for(raw, day)
        self.assertEqual([p.name for p in files], ['2026 10 28 00 00 00.txt', '2026 10 28 00 00 30.txt'])
        t, x, y = read_lemi_raw(files[1])
        self.assertEqual(len(t), 30)
        self.assertAlmostEqual(x[0], 35780.0)
        t0, x0, _ = read_lemi_raw(files[0])
        self.assertEqual(float(x0[0]), 35000.0)       # data PC tidak ditimpa

    def test_skips_sectors_without_valid_clock_or_magic(self):
        from .lemi_flash import read_flash
        folder = tempfile.mkdtemp(prefix='lemi-raw-')
        self.addCleanup(shutil.rmtree, folder, True)
        path = self._flash_file(folder, datetime.datetime(2000, 1, 1), 1)      # jam belum sinkron
        self.assertEqual(read_flash(path), {})
        with open(path, 'wb') as fh:
            fh.write(b'\xff' * 1024)
        self.assertEqual(read_flash(path), {})


class LemiDataAlertTest(TestCase):
    def setUp(self):
        from django.core.cache import cache
        self.raw = tempfile.mkdtemp(prefix='lemi-raw-')
        self.addCleanup(shutil.rmtree, self.raw, True)
        cache.delete('lemi-data-down-since')
        self.addCleanup(cache.delete, 'lemi-data-down-since')
        ctx = override_settings(LEMI_RAW_DIR=self.raw)
        ctx.enable()
        self.addCleanup(ctx.disable)

    def _sample(self, minutes_ago):
        t = datetime.datetime.now(UTC) - datetime.timedelta(minutes=minutes_ago)
        name = f'{t:%Y %m %d %H %M %S}.txt'
        with open(os.path.join(self.raw, name), 'w') as fh:
            fh.write(f'{t:%Y %m %d %H %M %S}  35780.00  -1820.00 -14790.00  29.00  28.00 P N FL  93\n')

    def test_alerts_once_when_data_stops_and_once_when_it_returns(self):
        from .tasks import check_lemi_data_task
        self._sample(90)
        with mock.patch('theme.telegram.send_telegram_message', return_value=True) as send:
            check_lemi_data_task()
            check_lemi_data_task()
        self.assertEqual(send.call_count, 1)
        self.assertIn('berhenti masuk', send.call_args[0][0])
        self._sample(5)
        with mock.patch('theme.telegram.send_telegram_message', return_value=True) as send:
            check_lemi_data_task()
            check_lemi_data_task()
        self.assertEqual(send.call_count, 1)
        self.assertIn('masuk lagi', send.call_args[0][0])

    def test_no_alert_when_data_is_fresh(self):
        from .tasks import check_lemi_data_task
        self._sample(5)
        with mock.patch('theme.telegram.send_telegram_message', return_value=True) as send:
            check_lemi_data_task()
        send.assert_not_called()
