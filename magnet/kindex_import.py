"""Impor indeks K dari berkas Excel pengolahan stasiun (grafik_kindeks <tahun>.xlsx).

Format per sheet (satu sheet per bulan, nama sheet = nama bulan Indonesia):
baris 5 dst: kolom A = tanggal, B..I = K(00-03)..K(21-24), J = indeks A, K = kriteria.
Hari tanpa data berisi teks (mis. "Lost Data") di kolom B atau kosong.
Tahun diambil dari argumen, bukan dari judul sheet (judul sheet bisa keliru).
Hanya tanggal sebelum KINDEX_FMI_START yang diimpor; mulai tanggal itu data diambil
dari citra FMI (lihat fmi_kindex.py).
"""
import calendar
import datetime

from openpyxl import load_workbook

from .fmi_kindex import fmi_start_date
from .models import KIndexDaily

MONTHS = {
    "januari": 1, "februari": 2, "maret": 3, "april": 4, "mei": 5, "juni": 6,
    "juli": 7, "agustus": 8, "september": 9, "oktober": 10, "november": 11, "desember": 12,
}


def import_kindex_workbook(fileobj, year, source_name="", months=None):
    """Impor sheet bulan dari workbook. Kembalikan (ringkasan per bulan, peringatan)."""
    wb = load_workbook(fileobj, data_only=True, read_only=True)
    summary, warnings = {}, []
    cutoff = fmi_start_date()
    skipped = 0
    for ws in wb.worksheets:
        month = MONTHS.get(ws.title.strip().lower())
        if not month or (months and month not in months):
            continue
        ndays = calendar.monthrange(year, month)[1]
        header = list(ws.iter_rows(min_row=2, max_row=2, max_col=11, values_only=True))[0]
        sheet_year = header[10]
        if isinstance(sheet_year, (int, float)) and int(sheet_year) != year:
            warnings.append(f"Sheet {ws.title}: tahun di judul sheet {int(sheet_year)}, diimpor sebagai {year}.")
        counts = {"data": 0, "kosong": 0}
        for row in ws.iter_rows(min_row=5, max_row=4 + ndays, max_col=11, values_only=True):
            day = row[0]
            if not isinstance(day, (int, float)) or not 1 <= int(day) <= ndays:
                continue
            ks = row[1:9]
            tanggal = datetime.date(year, month, int(day))
            if tanggal >= cutoff:
                skipped += 1
                continue
            if all(isinstance(k, (int, float)) and 0 <= k <= 9 for k in ks):
                values = [int(k) for k in ks]
                keterangan = ""
                counts["data"] += 1
            else:
                values = []
                keterangan = str(ks[0]).strip().title() if isinstance(ks[0], str) and ks[0].strip() else "Tidak ada data"
                counts["kosong"] += 1
            source_a = row[9] if isinstance(row[9], (int, float)) else None
            obj, _ = KIndexDaily.objects.get_or_create(tanggal=tanggal)
            obj.k_values = values
            obj.a_index_source = source_a
            obj.keterangan = keterangan
            obj.source_file = source_name
            obj.sumber = KIndexDaily.SUMBER_EXCEL
            obj.save()
        if counts["data"] or counts["kosong"]:
            summary[f"{year}-{month:02d}"] = counts
    if skipped:
        warnings.append(f"{skipped} hari mulai {cutoff:%d/%m/%Y} dilewati: data indeks K sejak tanggal itu diambil dari citra FMI.")
    return summary, warnings
