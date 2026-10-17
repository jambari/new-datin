from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings

from .models import KAHSeriesData, SPMSeriesData

KAH_CSV = (
    "stasiun,bulan,minggu,PH,DH,CA,MG,NA,K,NH4,CL,SO4,NO3,KT,AC,AL,CH,OX\n"
    "ANGKASAPURA,2026-01,01,5.20,4.4,0.067,0.031,0.181,0.181,0.034,0.355,0.240,0.147,0.098,10.5,0,8.9,0\n"
    "ANGKASAPURA,2026-01,02,4.64,9.7,0.136,0.052,0.312,0.001,0.001,0.526,0.266,0.079,0.188,29.1,0,40.2,0\n"
)
SPM_CSV = "stasiun,timestamp,SPM\nANGKASAPURA,2026-01-12,20.03\nANGKASAPURA,2026-01-18,22.62\n"
TOKEN = "test-token-123"
URL = "/kualitas-udara/api/import/"


def upload(name, text):
    return SimpleUploadedFile(name, text.encode(), content_type="text/csv")


@override_settings(KUALITAS_UDARA_API_TOKEN=TOKEN)
class ImportAPITest(TestCase):
    def post(self, token=TOKEN, **files):
        headers = {"HTTP_AUTHORIZATION": f"Token {token}"} if token else {}
        return self.client.post(URL, files, **headers)

    def test_requires_valid_token(self):
        self.assertEqual(self.post(token=None, kah=upload("k.csv", KAH_CSV)).status_code, 401)
        self.assertEqual(self.post(token="salah", kah=upload("k.csv", KAH_CSV)).status_code, 401)
        self.assertEqual(KAHSeriesData.objects.count(), 0)

    @override_settings(KUALITAS_UDARA_API_TOKEN="")
    def test_disabled_without_token_setting(self):
        self.assertEqual(self.post(kah=upload("k.csv", KAH_CSV)).status_code, 503)

    def test_get_not_allowed(self):
        self.assertEqual(self.client.get(URL, HTTP_AUTHORIZATION=f"Token {TOKEN}").status_code, 405)

    def test_imports_and_skips_existing_without_overwriting(self):
        r = self.post(kah=upload("k.csv", KAH_CSV), spm=upload("s.csv", SPM_CSV))
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["results"]["kah"]["created"], 2)
        self.assertEqual(r.json()["results"]["spm"]["created"], 2)

        KAHSeriesData.objects.filter(bulan="2026-01", minggu="01").update(ph=9.99)
        newer = KAH_CSV + "ANGKASAPURA,2026-01,03,4.37,22.2,0.393,0.094,0.353,0.001,0.001,1.11,0.338,0.006,0.487,0,0,98.2,0\n"
        r = self.post(kah=upload("k.csv", newer))
        self.assertEqual(r.json()["results"]["kah"], {"created": 1, "updated": 0, "skipped": 2, "invalid": 0})
        self.assertEqual(KAHSeriesData.objects.get(bulan="2026-01", minggu="01").ph, 9.99)
        self.assertEqual(SPMSeriesData.objects.count(), 2)

    def test_rejects_wrong_csv(self):
        r = self.post(kah=upload("k.csv", "a,b,c\n1,2,3\n"))
        self.assertEqual(r.status_code, 400)
        self.assertIn("tidak memuat kolom", r.json()["error"])
        self.assertEqual(self.post().status_code, 400)


class ImportCommandTest(TestCase):
    def test_skip_existing_flag(self):
        import tempfile, os
        with tempfile.TemporaryDirectory() as d:
            kah, spm = os.path.join(d, "k.csv"), os.path.join(d, "s.csv")
            open(kah, "w").write(KAH_CSV); open(spm, "w").write(SPM_CSV)
            call_command("import_kualitas_udara", kah=kah, spm=spm, stdout=open(os.devnull, "w"))
            KAHSeriesData.objects.update(ph=1.0)
            call_command("import_kualitas_udara", kah=kah, spm=spm, skip_existing=True, stdout=open(os.devnull, "w"))
            self.assertTrue(all(x == 1.0 for x in KAHSeriesData.objects.values_list("ph", flat=True)))
            call_command("import_kualitas_udara", kah=kah, spm=spm, stdout=open(os.devnull, "w"))
            self.assertEqual(KAHSeriesData.objects.get(minggu="01").ph, 5.2)


class FilterTest(TestCase):
    def setUp(self):
        self.client.force_login(get_user_model().objects.create_user("t", password="x"))
        for b in ("2025-12", "2026-01", "2026-03", "2026-04"):
            KAHSeriesData.objects.create(stasiun="ANGKASAPURA", bulan=b, minggu="01", ph=5.0)

    def months(self, url):
        return sorted(row["obj"].bulan for row in self.client.get(url).context["rows"])

    def test_month_year_dropdowns_and_legacy_params(self):
        self.assertEqual(self.months("/kualitas-udara/kah/?start_bulan=01&start_tahun=2026&end_bulan=03&end_tahun=2026"),
                         ["2026-01", "2026-03"])
        self.assertEqual(self.months("/kualitas-udara/kah/?end_tahun=2025"), ["2025-12"])
        self.assertEqual(self.months("/kualitas-udara/kah/?start=2026-01&end=2026-03"), ["2026-01", "2026-03"])
        html = self.client.get("/kualitas-udara/kah/").content.decode()
        self.assertNotIn('type="month"', html)
        self.assertNotIn('name="stasiun"', html)


class EmptyRowAndParameterInfoTest(TestCase):
    def setUp(self):
        self.client.force_login(get_user_model().objects.create_user("t", password="x"))
        self.empty = KAHSeriesData.objects.create(
            stasiun="ANGKASAPURA", bulan="2026-01", minggu="05",
            **{f: 0.0 for f in ("ph", "dh", "ca", "mg", "na", "k", "nh4", "cl", "so4", "no3", "kt", "ac", "al", "ch", "ox")},
        )
        self.acid = KAHSeriesData.objects.create(stasiun="ANGKASAPURA", bulan="2026-01", minggu="01", ph=4.5, ch=12.0)

    def test_all_zero_row_is_no_sample_not_exceedance(self):
        self.assertTrue(self.empty.is_empty)
        self.assertFalse(self.acid.is_empty)
        self.assertEqual(self.empty.get_remark("PH")["status"], "nosample")
        self.assertEqual(self.empty.get_exceeded_fields(), [])
        self.assertEqual(self.empty.get_remark_summary(), "Tidak ada sampel")
        self.assertEqual(self.acid.get_remark("PH")["status"], "exceed")

    def test_list_page_shows_no_sample_and_parameter_legend(self):
        html = self.client.get("/kualitas-udara/kah/").content.decode()
        self.assertIn("Tidak ada sampel", html)
        self.assertEqual(html.count("Melebihi ambang"), 1)  # only the real pH 4.5 sample
        self.assertIn("Keterangan parameter KAH", html)
        self.assertIn("Kekerasan total", html)
        self.assertIn("µeq/L", html)

    def test_export_uses_units_and_no_sample_remark(self):
        import io
        from openpyxl import load_workbook
        r = self.client.get("/kualitas-udara/kah/export/excel/")
        ws = load_workbook(io.BytesIO(r.content)).active
        header = [c.value for c in ws[1]]
        self.assertIn("AC (µeq/L)", header)
        self.assertIn("CH (mm)", header)
        remarks = [row[-1] for row in ws.iter_rows(min_row=2, values_only=True)]
        self.assertIn("Tidak ada sampel", remarks)
