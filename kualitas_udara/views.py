import hmac
import io
import logging
import re
from datetime import datetime
from urllib.parse import urlencode

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import EmptyPage, PageNotAnInteger, Paginator
from django.db import transaction
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from openpyxl import Workbook
from openpyxl.styles import Font
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_POST
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .forms import KAHSeriesDataForm, SPMSeriesDataForm
from .importer import CSVFormatError, import_kah, import_spm
from .models import (
    KAH_FIELD_LABELS,
    KAH_FIELDS,
    KAH_THRESHOLDS,
    KAHSeriesData,
    SPM_THRESHOLD,
    SPM_THRESHOLD_LABEL,
    SPMSeriesData,
)

logger = logging.getLogger(__name__)

PER_PAGE_OPTIONS = (10, 25, 50, 100)
DEFAULT_PER_PAGE = 10

MONTH_CHOICES = [
    ("01", "Januari"), ("02", "Februari"), ("03", "Maret"), ("04", "April"),
    ("05", "Mei"), ("06", "Juni"), ("07", "Juli"), ("08", "Agustus"),
    ("09", "September"), ("10", "Oktober"), ("11", "November"), ("12", "Desember"),
]


# ---------------------------------------------------------------------------
# Helper umum
# ---------------------------------------------------------------------------
def get_per_page_value(request):
    try:
        value = int(request.GET.get("per_page", DEFAULT_PER_PAGE))
    except (TypeError, ValueError):
        return DEFAULT_PER_PAGE
    return value if value in PER_PAGE_OPTIONS else DEFAULT_PER_PAGE


def build_preserved_params(request, exclude=None):
    exclude = exclude or {"page"}
    params = {}
    for key, value in request.GET.items():
        if key not in exclude and value:
            params[key] = value
    return params


def build_filter_query(request, exclude=None):
    return urlencode(build_preserved_params(request, exclude))


def paginate_queryset(request, queryset, per_page=DEFAULT_PER_PAGE):
    paginator = Paginator(queryset, per_page)
    page = request.GET.get("page")
    try:
        return paginator.page(page)
    except PageNotAnInteger:
        return paginator.page(1)
    except EmptyPage:
        return paginator.page(paginator.num_pages)


def get_page_range(page_obj):
    return page_obj.paginator.get_elided_page_range(
        page_obj.number, on_each_side=2, on_ends=2
    )


# ---------------------------------------------------------------------------
# Filter queryset
# ---------------------------------------------------------------------------
def get_month_param(request, prefix, default_month):
    """Baca filter bulan sebagai 'YYYY-MM'.

    Dari dropdown <prefix>_bulan + <prefix>_tahun (berfungsi di semua browser,
    termasuk Firefox yang tidak mendukung <input type="month">), atau dari
    parameter lama <prefix>=YYYY-MM.
    """
    year = request.GET.get(f"{prefix}_tahun", "").strip()
    month = request.GET.get(f"{prefix}_bulan", "").strip() or default_month
    if re.fullmatch(r"\d{4}", year) and re.fullmatch(r"\d{2}", month):
        return f"{year}-{month}"
    raw = request.GET.get(prefix, "").strip()
    return raw if re.fullmatch(r"\d{4}-\d{2}", raw) else ""


def get_kah_year_options():
    years = {bulan[:4] for bulan in KAHSeriesData.objects.values_list("bulan", flat=True).distinct()}
    return sorted((y for y in years if y.isdigit()), reverse=True)


def filter_kah(request):
    queryset = KAHSeriesData.objects.all().order_by("-bulan", "-minggu", "-id")
    start = get_month_param(request, "start", "01")
    end = get_month_param(request, "end", "12")
    if start:
        queryset = queryset.filter(bulan__gte=start)
    if end:
        queryset = queryset.filter(bulan__lte=end)
    return queryset, start, end


def filter_spm(request):
    queryset = SPMSeriesData.objects.all().order_by("-timestamp", "-id")
    start = request.GET.get("start", "").strip()
    end = request.GET.get("end", "").strip()
    if start:
        queryset = queryset.filter(timestamp__gte=start)
    if end:
        queryset = queryset.filter(timestamp__lte=end)
    return queryset, start, end


# ---------------------------------------------------------------------------
# Penyusunan baris tabel (nilai + remark per kolom)
# ---------------------------------------------------------------------------
def build_kah_rows(records):
    rows = []
    for record in records:
        cells = {}
        for field_key in KAH_FIELDS:
            cells[field_key] = {
                "value": getattr(record, field_key.lower()),
                "remark": record.get_remark(field_key),
            }
        rows.append(
            {
                "obj": record,
                "cells": cells,
                "remark_summary": record.get_remark_summary(),
            }
        )
    return rows


def build_kah_field_defs():
    defs = []
    for field_key in KAH_FIELDS:
        spec = KAH_THRESHOLDS.get(field_key)
        defs.append(
            {
                "key": field_key,
                "label": KAH_FIELD_LABELS[field_key],
                "unit": spec["unit"] if spec else "",
                "threshold_label": spec["label"] if spec else "",
            }
        )
    return defs


def build_spm_rows(records):
    return [
        {
            "obj": record,
            "value": record.spm,
            "remark": record.get_remark(),
            "remark_summary": record.get_remark_summary(),
        }
        for record in records
    ]


# ---------------------------------------------------------------------------
# KAH (Kimia Air Hujan)
# ---------------------------------------------------------------------------
@login_required
def kah_list(request):
    queryset, start, end = filter_kah(request)
    per_page = get_per_page_value(request)
    page_obj = paginate_queryset(request, queryset, per_page=per_page)
    rows = build_kah_rows(page_obj.object_list)

    context = {
        "title": "Data Kimia Air Hujan (KAH)",
        "rows": rows,
        "fields": build_kah_field_defs(),
        "page_obj": page_obj,
        "page_range": get_page_range(page_obj),
        "per_page": per_page,
        "per_page_options": PER_PAGE_OPTIONS,
        "preserved_params": build_preserved_params(request, exclude={"page"}),
        "filter_query": build_filter_query(request, exclude={"page"}),
        "f_start_bulan": start[5:7],
        "f_start_tahun": start[:4],
        "f_end_bulan": end[5:7],
        "f_end_tahun": end[:4],
        "month_choices": MONTH_CHOICES,
        "year_options": get_kah_year_options(),
        "has_filters": any([start, end]),
        "threshold_notes": {
            key: KAH_THRESHOLDS[key]["label"]
            for key in KAH_THRESHOLDS
        },
    }
    return render(request, "kualitas_udara/kah_series_list.html", context)


@login_required
def kah_create(request):
    if request.method == "POST":
        form = KAHSeriesDataForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Data KAH berhasil ditambahkan.")
            return redirect("kualitas_udara:kah_list")
    else:
        form = KAHSeriesDataForm()
    return render(
        request,
        "kualitas_udara/form_kah.html",
        {"form": form, "title": "Tambah Data KAH"},
    )


@login_required
def kah_update(request, pk):
    record = get_object_or_404(KAHSeriesData, pk=pk)
    if request.method == "POST":
        form = KAHSeriesDataForm(request.POST, instance=record)
        if form.is_valid():
            form.save()
            messages.success(request, "Data KAH berhasil diperbarui.")
            return redirect("kualitas_udara:kah_list")
    else:
        form = KAHSeriesDataForm(instance=record)
    return render(
        request,
        "kualitas_udara/form_kah.html",
        {"form": form, "title": "Edit Data KAH", "record": record},
    )


@login_required
def kah_delete(request, pk):
    record = get_object_or_404(KAHSeriesData, pk=pk)
    if request.method == "POST":
        record.delete()
        messages.success(request, "Data KAH berhasil dihapus.")
        return redirect("kualitas_udara:kah_list")
    return render(
        request,
        "kualitas_udara/confirm_delete.html",
        {"object": record, "type_name": "Data KAH", "back_url": "kualitas_udara:kah_list"},
    )


@login_required
def kah_export(request, export_format):
    queryset, _start, _end = filter_kah(request)

    headers = ["Stasiun", "Bulan", "Minggu"] + [
        f"{KAH_FIELD_LABELS[f]}{' (' + KAH_THRESHOLDS[f]['unit'] + ')' if f in KAH_THRESHOLDS and KAH_THRESHOLDS[f]['unit'] else ''}"
        for f in KAH_FIELDS
    ] + ["Remark"]

    def iter_rows():
        for record in queryset:
            yield [record.stasiun, record.bulan, record.minggu] + [
                getattr(record, field_key.lower()) for field_key in KAH_FIELDS
            ] + [record.get_remark_summary()]

    if export_format == "excel":
        return export_table_excel(
            "KAH",
            headers,
            iter_rows,
            "kah_series_data",
        )
    return export_table_pdf(
        "Data Kimia Air Hujan (KAH) - Kualitas Udara",
        headers,
        iter_rows,
        "kah_series_data",
    )


# ---------------------------------------------------------------------------
# SPM (Suspended Particulate Matter)
# ---------------------------------------------------------------------------
@login_required
def spm_list(request):
    queryset, start, end = filter_spm(request)
    per_page = get_per_page_value(request)
    page_obj = paginate_queryset(request, queryset, per_page=per_page)
    rows = build_spm_rows(page_obj.object_list)

    context = {
        "title": "Data SPM (Partikulat)",
        "rows": rows,
        "page_obj": page_obj,
        "page_range": get_page_range(page_obj),
        "per_page": per_page,
        "per_page_options": PER_PAGE_OPTIONS,
        "preserved_params": build_preserved_params(request, exclude={"page"}),
        "filter_query": build_filter_query(request, exclude={"page"}),
        "f_start": start,
        "f_end": end,
        "has_filters": any([start, end]),
        "spm_threshold": SPM_THRESHOLD,
        "spm_threshold_label": SPM_THRESHOLD_LABEL,
    }
    return render(request, "kualitas_udara/spm_series_list.html", context)


@login_required
def spm_create(request):
    if request.method == "POST":
        form = SPMSeriesDataForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Data SPM berhasil ditambahkan.")
            return redirect("kualitas_udara:spm_list")
    else:
        form = SPMSeriesDataForm()
    return render(
        request,
        "kualitas_udara/form_spm.html",
        {"form": form, "title": "Tambah Data SPM"},
    )


@login_required
def spm_update(request, pk):
    record = get_object_or_404(SPMSeriesData, pk=pk)
    if request.method == "POST":
        form = SPMSeriesDataForm(request.POST, instance=record)
        if form.is_valid():
            form.save()
            messages.success(request, "Data SPM berhasil diperbarui.")
            return redirect("kualitas_udara:spm_list")
    else:
        form = SPMSeriesDataForm(instance=record)
    return render(
        request,
        "kualitas_udara/form_spm.html",
        {"form": form, "title": "Edit Data SPM", "record": record},
    )


@login_required
def spm_delete(request, pk):
    record = get_object_or_404(SPMSeriesData, pk=pk)
    if request.method == "POST":
        record.delete()
        messages.success(request, "Data SPM berhasil dihapus.")
        return redirect("kualitas_udara:spm_list")
    return render(
        request,
        "kualitas_udara/confirm_delete.html",
        {"object": record, "type_name": "Data SPM", "back_url": "kualitas_udara:spm_list"},
    )


@login_required
def spm_export(request, export_format):
    queryset, _start, _end = filter_spm(request)
    headers = ["Stasiun", "Tanggal", "SPM (µg/m³)", "Remark"]

    def iter_rows():
        for record in queryset:
            yield [
                record.stasiun,
                record.timestamp.strftime("%Y-%m-%d") if record.timestamp else "",
                record.spm,
                record.get_remark_summary(),
            ]

    if export_format == "excel":
        return export_table_excel("SPM", headers, iter_rows, "spm_series_data")
    return export_table_pdf(
        "Data SPM (Partikulat) - Kualitas Udara",
        headers,
        iter_rows,
        "spm_series_data",
    )


# ---------------------------------------------------------------------------
# Export helper (Excel & PDF)
# ---------------------------------------------------------------------------
def export_table_excel(sheet_title, headers, row_iterator, filename_prefix):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row in row_iterator():
        sheet.append(["" if value is None else value for value in row])

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = f"{filename_prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _pdf_safe(value):
    """Ganti karakter subskrip unicode agar aman untuk font standar reportlab."""
    if isinstance(value, str):
        return value.replace("₄", "4").replace("₃", "3")
    return value


def export_table_pdf(doc_title, headers, row_iterator, filename_prefix):
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=landscape(A4), leftMargin=1 * cm, rightMargin=1 * cm
    )
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(doc_title, styles["Title"]),
        Spacer(1, 12),
    ]

    table_data = [[_pdf_safe(header) for header in headers]]
    for row in row_iterator():
        table_data.append([_pdf_safe("" if value is None else value) for value in row])

    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 6.5),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f9fafb")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    elements.append(table)
    document.build(elements)

    filename = f"{filename_prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# ---------------------------------------------------------------------------
# API impor (dipanggil cron di scclient setelah menarik CSV dari
# Database Kualitas Udara BMKG). Hanya menambah data baru; data yang sudah
# ada tidak pernah ditimpa.
# ---------------------------------------------------------------------------
API_MAX_UPLOAD_BYTES = 5 * 1024 * 1024
API_IMPORTERS = (("kah", import_kah), ("spm", import_spm))


def _client_ip(request):
    return request.headers.get("CF-Connecting-IP") or request.META.get("REMOTE_ADDR", "")


@csrf_exempt
@require_POST
def api_import(request):
    token = getattr(settings, "KUALITAS_UDARA_API_TOKEN", "")
    if not token:
        return JsonResponse({"ok": False, "error": "API impor kualitas udara belum diaktifkan."}, status=503)

    auth = request.headers.get("Authorization", "")
    given = auth[len("Token "):].strip() if auth.startswith("Token ") else ""
    if not given or not hmac.compare_digest(given.encode(), token.encode()):
        logger.warning("Kualitas udara API: token tidak valid dari %s", _client_ip(request))
        return JsonResponse({"ok": False, "error": "Token tidak valid."}, status=401)

    uploads = {key: request.FILES.get(key) for key, _ in API_IMPORTERS}
    if not any(uploads.values()):
        return JsonResponse({"ok": False, "error": "Kirim file CSV pada field 'kah' dan/atau 'spm'."}, status=400)

    results = {}
    for key, importer in API_IMPORTERS:
        upload = uploads[key]
        if upload is None:
            continue
        if upload.size > API_MAX_UPLOAD_BYTES:
            return JsonResponse({"ok": False, "error": f"File {key} terlalu besar."}, status=413)
        try:
            lines = upload.read().decode("utf-8-sig").splitlines()
            with transaction.atomic():
                results[key] = importer(lines, skip_existing=True)
        except UnicodeDecodeError:
            return JsonResponse({"ok": False, "error": f"File {key} bukan teks UTF-8."}, status=400)
        except CSVFormatError as exc:
            return JsonResponse({"ok": False, "error": str(exc), "results": results}, status=400)

    logger.info("Kualitas udara API: impor dari %s: %s", _client_ip(request), results)
    return JsonResponse({"ok": True, "results": results})
