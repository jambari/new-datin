import datetime
import io
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import Workbook
from PIL import Image

from kualitas_udara.models import KAHSeriesData, SPMSeriesData
from lightning.models import DailyStrikeSummary, Strike
from magnet.kindex_import import import_kindex_workbook
from magnet.models import KIndexDaily, MagneticObservation
from repository.models import EventResponseSpectrum, ShakemapEvent, Station

UTC = datetime.timezone.utc


def fake_tile(z, x, y):
    return Image.new("RGB", (256, 256), (200, 220, 240))


def report_url(jenis, start=(1, 2026), end=(3, 2026), **extra):
    params = {"jenis": jenis, "start_bulan": start[0], "start_tahun": start[1], "end_bulan": end[0],
              "end_tahun": end[1], "nip": "199202062012101001",
              **extra}
    return "/reporting/laporan/?" + "&".join(f"{k}={v}" for k, v in params.items())


def kindex_workbook(sheet_year=2026):
    wb = Workbook()
    ws = wb.active
    ws.title = "Januari"
    ws["I2"], ws["J2"], ws["K2"] = "Bulan :", "Januari", sheet_year
    ws.append([])
    ws["A4"] = "DAY"
    for day in range(1, 32):
        row = 4 + day
        ws.cell(row, 1, day)
        if day == 20:
            ws.cell(row, 2, "Lost Data")
        else:
            ks = [8, 9, 4, 3, 2, 2, 1, 1] if day == 19 else [2, 3, 2, 1, 1, 2, 2, 3]
            for j, k in enumerate(ks):
                ws.cell(row, 2 + j, k)
            ws.cell(row, 10, 999 if day == 5 else None)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ReportingTest(TestCase):
    def setUp(self):
        self.client.force_login(get_user_model().objects.create_user("t", password="x", first_name="Jambari"))

    def test_requires_login(self):
        self.client.logout()
        self.assertEqual(self.client.get("/reporting/").status_code, 302)
        self.assertEqual(self.client.get(report_url("kualitas_udara")).status_code, 302)

    def test_index_page(self):
        r = self.client.get("/reporting/")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Laporan Bukti SKP")  # sidebar link
        self.assertContains(r, "Magnetbumi &amp; Listrik Udara")

    def test_invalid_period_redirects_with_message(self):
        r = self.client.get(report_url("kualitas_udara", start=(3, 2026), end=(1, 2026)), follow=True)
        self.assertContains(r, "Bulan akhir tidak boleh sebelum bulan awal")
        r = self.client.get(report_url("kualitas_udara", start=(1, 2025), end=(3, 2026)), follow=True)
        self.assertContains(r, "maksimal 12 bulan")

    def test_kualitas_udara_report(self):
        zeros = {f: 0.0 for f in ("ph", "dh", "ca", "mg", "na", "k", "nh4", "cl", "so4", "no3", "kt", "ac", "al", "ch", "ox")}
        KAHSeriesData.objects.create(stasiun="ANGKASAPURA", bulan="2026-01", minggu="05", **zeros)
        for i, ph in enumerate([4.5, 5.2, 4.9, 6.0], 1):
            KAHSeriesData.objects.create(stasiun="ANGKASAPURA", bulan="2026-02", minggu=f"0{i}", ph=ph, dh=8.0,
                                         ca=0.1, mg=0.1, na=0.4, cl=0.7, so4=0.3, kt=0.2, ac=10.0, ch=20.0)
        KAHSeriesData.objects.create(stasiun="ANGKASAPURA", bulan="2026-03", minggu="01", ph=7.3, dh=160.0,
                                     ca=5.0, mg=23.0, kt=28.0, ac=0.0)
        for d in (6, 12, 18, 30):
            SPMSeriesData.objects.create(stasiun="ANGKASAPURA", timestamp=datetime.date(2026, 1, d), spm=20.0 + d)
        r = self.client.get(report_url("kualitas_udara"))
        self.assertEqual(r.status_code, 200)
        c = r.context
        self.assertEqual(len(c["valid"]), 5)
        self.assertEqual(len(c["empty"]), 1)
        self.assertEqual(len(c["verify"]), 1)  # March outlier
        self.assertEqual(c["acid"], 3)
        self.assertEqual([d.day for d in c["spm_missing"]], [24])
        self.assertContains(r, "Tersedianya data Kualitas Udara yang berkualitas")
        self.assertContains(r, "199202062012101001")
        self.assertContains(r, "data:image/png;base64,")

    @mock.patch("reporting.osm._tile", side_effect=fake_tile)
    def test_magnet_petir_report(self, _tile):
        DailyStrikeSummary.objects.create(summary_date=datetime.date(2026, 1, 2), cg_plus_count=10,
                                          cg_minus_count=5, ic_count=3, total_count=18)
        for i in range(60):
            Strike.objects.create(epoch_ms=i, timestamp=datetime.datetime(2026, 1, 2, 3, i % 60, tzinfo=UTC),
                                  latitude=-2.3 - i * 0.005, longitude=140.5 + i * 0.005, strike_type=i % 3)
        for i in range(10):  # duplicate rows (same file ingested twice) count once
            Strike.objects.create(epoch_ms=i, timestamp=datetime.datetime(2026, 1, 2, 3, i % 60, tzinfo=UTC),
                                  latitude=-2.3 - i * 0.005, longitude=140.5 + i * 0.005, strike_type=i % 3)
        Strike.objects.create(epoch_ms=999, timestamp=datetime.datetime(2026, 1, 2, 4, 0, tzinfo=UTC),
                              latitude=-5.0, longitude=138.0, strike_type=0)  # outside the summary box
        MagneticObservation.objects.create(observation_date=datetime.date(2026, 3, 25), observer="Juan",
                                           session="Session 1", declination=2.79, inclination=-22.1,
                                           total_intensity=39263.08)
        MagneticObservation.objects.create(observation_date=datetime.date(2026, 1, 16), observer="Jambari",
                                           session="Session 1",
                                           deklinasi_readings={"WU": {"hh": "0", "mm": "47", "ss": "1", "circ": "303.5"}},
                                           inklinasi_readings={k: {"ftotal": "39273.0"} for k in ("NU", "SD", "ND", "SU")})
        import_kindex_workbook(io.BytesIO(kindex_workbook()), 2026, "grafik_kindeks 2026.xlsx")
        r = self.client.get(report_url("magnet_petir"))
        self.assertEqual(r.status_code, 200)
        c = r.context
        self.assertEqual(c["petir_total"], 18)
        self.assertEqual(c["petir"][0]["map_points"], 60)
        self.assertIsNotNone(c["petir"][0]["map"])
        self.assertTrue(c["petir"][0]["map_mismatch"])  # 60 unique points vs 18 in the summary
        self.assertContains(r, "berbeda lebih dari 5%")
        self.assertEqual(len(c["magnet_reduced"]), 1)
        self.assertEqual(len(c["magnet_dim"]), 1)
        self.assertEqual(c["magnet_empty_months"], ["Februari 2026"])
        self.assertEqual(c["kdays"], 30)
        self.assertEqual([o.tanggal.day for o, _crit in c["k_storms"]], [19])
        self.assertContains(r, "OpenStreetMap")

    def test_kindex_import(self):
        summary, warnings = import_kindex_workbook(io.BytesIO(kindex_workbook(sheet_year=2025)), 2026, "k.xlsx")
        self.assertTrue(any("tahun di judul sheet 2025" in w for w in warnings))
        self.assertEqual(summary["2026-01"]["data"], 30)
        lost = KIndexDaily.objects.get(tanggal=datetime.date(2026, 1, 20))
        self.assertFalse(lost.has_data)
        self.assertEqual(lost.keterangan, "Lost Data")
        storm = KIndexDaily.objects.get(tanggal=datetime.date(2026, 1, 19))
        self.assertEqual(storm.a_index, (240 + 400 + 27 + 15 + 7 + 7 + 3 + 3) / 8)
        self.assertEqual(storm.criteria, "Badai Menengah")
        self.assertEqual(KIndexDaily.objects.get(tanggal=datetime.date(2026, 1, 5)).a_index_source, 999)

    def test_kindex_page(self):
        import_kindex_workbook(io.BytesIO(kindex_workbook()), 2026, "grafik_kindeks 2026.xlsx")
        r = self.client.get("/reporting/indeks-k/?start_bulan=1&start_tahun=2026&end_bulan=2&end_tahun=2026")
        self.assertEqual(r.status_code, 200)
        c = r.context
        self.assertEqual(c["kdays"], 30)
        self.assertEqual([k["n"] for k in c["kmonths"]], [30, 0])
        self.assertEqual([o.tanggal.day for o, _crit in c["k_storms"]], [19])
        self.assertContains(r, "data:image/png;base64,", count=1)  # chart only for the month with data
        self.assertContains(r, "Lost Data")
        self.assertContains(r, "Badai Menengah")
        self.assertContains(r, "Belum ada data indeks K untuk Februari 2026")

    def test_kindex_page_shows_bartington_absolute_values(self):
        def dms(d, m=0):
            return {"deg": str(d), "min": str(m), "sec": "0", "hh": "1", "mm": "0"}
        MagneticObservation.objects.create(
            observation_date=datetime.date(2026, 1, 21), observer="Achmad", session="Session 1",
            deklinasi_readings={"WU": dms(182, 48), "ED": dms(182, 50), "WD": dms(182, 52), "EU": dms(182, 50), "_cr_awal": True},
            inklinasi_readings={k: {**dms(v), "ftotal": "39270"} for k, v in
                                (("NU", 202), ("SD", 338), ("ND", 158), ("SU", 22))})
        # DIM (MinGeo) dan data impor lama bukan Bartington: tidak ditampilkan
        MagneticObservation.objects.create(observation_date=datetime.date(2026, 1, 16), observer="Jambari",
                                           session="Session 1",
                                           deklinasi_readings={"WU": {"hh": "0", "mm": "47", "ss": "1", "circ": "303.5"}})
        MagneticObservation.objects.create(observation_date=datetime.date(2026, 1, 25), observer="Juan",
                                           session="Session 1", declination=2.79, inclination=-22.1,
                                           total_intensity=39263.08)
        r = self.client.get("/reporting/indeks-k/?start_bulan=1&start_tahun=2026&end_bulan=2&end_tahun=2026")
        c = r.context
        self.assertEqual(c["abs_total"], 1)
        self.assertEqual([len(a["rows"]) for a in c["abs_months"]], [1, 0])
        self.assertEqual(c["abs_months"][0]["rows"][0].observer, "Achmad")
        self.assertIsNotNone(c["abs_months"][0]["chart"])
        self.assertContains(r, "Nilai absolut – Januari 2026")
        self.assertContains(r, "Tidak ada observasi absolut Bartington pada Februari 2026")
        self.assertNotContains(r, "Juan")

    def test_kindex_page_defaults_to_last_month(self):
        r = self.client.get("/reporting/indeks-k/")
        self.assertEqual(r.status_code, 200)
        last = datetime.date.today().replace(day=1) - datetime.timedelta(days=1)
        self.assertEqual((r.context["period"].start.year, r.context["period"].start.month), (last.year, last.month))

    def test_kindex_page_rejects_bad_period(self):
        r = self.client.get("/reporting/indeks-k/?start_bulan=3&start_tahun=2026&end_bulan=1&end_tahun=2026")
        self.assertEqual(r.status_code, 200)
        self.assertNotIn("period", r.context)
        self.assertContains(r, "Bulan akhir tidak boleh sebelum bulan awal.")

    def test_kindex_page_requires_login(self):
        self.client.logout()
        self.assertEqual(self.client.get("/reporting/indeks-k/").status_code, 302)

    def test_no_kindex_upload_on_reporting_page(self):
        r = self.client.get("/reporting/")
        self.assertNotContains(r, 'type="file"')
        self.assertNotContains(r, "Unggah")
        self.assertContains(r, "Indeks K tersimpan")
        self.assertEqual(self.client.post("/reporting/indeks-k/upload/").status_code, 404)

    def test_percepatan_tanah_report(self):
        from repository.models import FeltEarthquake
        Station.objects.create(code="JAY", network="IA", latitude=-2.51, longitude=140.70)
        Station.objects.create(code="SWI", network="IA", latitude=-2.60, longitude=140.30)
        # Record lama: id XML panjang, event_time = jam WIB tertulis UTC.
        ShakemapEvent.objects.create(event_id="20260114064812_32_-2600_140600_20260114064812", latitude=-2.6,
                                     longitude=140.6, magnitude=3.2, depth=10,
                                     event_time=datetime.datetime(2026, 1, 14, 6, 48, 12, tzinfo=UTC),
                                     location_string="10 km Jayapura", shakemap_image="shakemaps/old.jpg")
        # Record dibuat otomatis saat push (metadata 0) -> metadata dari gempa dirasakan.
        ShakemapEvent.objects.create(event_id="20260401050000", latitude=0, longitude=0, magnitude=0, depth=0,
                                     event_time=datetime.datetime(2026, 4, 1, 5, 0, tzinfo=UTC),
                                     location_string="Pushed from jaygenerator (auto-created)")
        FeltEarthquake.objects.create(event_datetime=datetime.datetime(2026, 3, 31, 22, 0, 2, tzinfo=UTC),
                                      latitude=-2.55, longitude=140.65, magnitude=4.1, depth_km=12,
                                      wilayah="Dekat Jayapura")

        def spec(scale):
            return [{"T": 0.0, "Sa": 0.02 * scale}, {"T": 0.3, "Sa": 0.04 * scale},
                    {"T": 1.0, "Sa": 0.01 * scale}, {"T": 3.0, "Sa": 0.001 * scale}]
        for comp in ("HNE", "HNN", "HNZ"):
            EventResponseSpectrum.objects.create(event_id="20260114064812", station_code="JAY", component=comp, spectrum=spec(1))
            EventResponseSpectrum.objects.create(event_id="20260114064812", station_code="SWI", component=comp,
                                                 spectrum=spec(2000 if comp == "HNE" else 1))
        # 1 Apr 05:00 WIB = 31 Mar 22:00 UTC -> masuk Januari–Maret.
        EventResponseSpectrum.objects.create(event_id="20260401050000", station_code="JAY", component="HNE", spectrum=spec(1))
        # 1 Apr 08:00 WIB = 1 Apr 01:00 UTC -> di luar periode.
        EventResponseSpectrum.objects.create(event_id="20260401080000", station_code="JAY", component="HNE", spectrum=spec(1))
        r = self.client.get(report_url("percepatan_tanah"))
        self.assertEqual(r.status_code, 200)
        c = r.context
        self.assertEqual([e["event_id"] for e in c["events"]], ["20260114064812", "20260401050000"])
        first, last = c["events"]
        self.assertEqual(first["time"], datetime.datetime(2026, 1, 13, 23, 48, 12, tzinfo=UTC))  # WIB -> UTC
        self.assertEqual(first["image"].name, "shakemaps/old.jpg")  # id lama tetap ditemukan
        self.assertEqual(first["nearest"]["station"], "JAY")
        self.assertAlmostEqual(first["pga"]["pga"], 0.02)
        self.assertEqual(c["n_bad"], 3)  # SWI spike + komponen lainnya
        self.assertEqual(last["mag"], 4.1)  # dari gempa dirasakan
        self.assertEqual(last["location"], "Dekat Jayapura")
        self.assertIsNone(last["image"])
        self.assertContains(r, "nilai &gt; 3 g")


class OfficerIdentityTest(TestCase):
    def test_name_and_nip_come_from_logged_in_pegawai(self):
        from jadwal.models import Pegawai
        user = get_user_model().objects.create_user("jambari", password="x", first_name="Jambari")
        Pegawai.objects.create(user=user, nama="Jambari, S.Tr, M.Si", nip="199202062012101001",
                               pangkat="Penata", jabatan="PMG")
        self.client.force_login(user)
        index = self.client.get("/reporting/")
        self.assertContains(index, "Jambari, S.Tr, M.Si")
        self.assertNotContains(index, 'name="nama"')
        self.assertNotContains(index, "atasan")
        r = self.client.get(report_url("kualitas_udara", nip="000", nama="Orang Lain"))
        self.assertEqual(r.context["identity"]["nama"], "Jambari, S.Tr, M.Si")
        self.assertEqual(r.context["identity"]["nip"], "199202062012101001")
        self.assertNotContains(r, "Orang Lain")
        self.assertNotContains(r, "Mengetahui")
        self.assertNotContains(r, "Yang Membuat")
        self.assertContains(r, "Stasiun Geofisika Jayapura")
        self.assertNotContains(index, 'name="unit_kerja"')
        self.assertNotContains(index, 'name="tempat_tanggal"')

    def test_user_without_pegawai_uses_account_name_and_manual_nip(self):
        user = get_user_model().objects.create_user("tamu", password="x", first_name="Budi", last_name="Santoso")
        self.client.force_login(user)
        self.assertContains(self.client.get("/reporting/"), 'name="nip"')
        r = self.client.get(report_url("kualitas_udara", nip="123456"))
        self.assertEqual(r.context["identity"]["nama"], "Budi Santoso")
        self.assertEqual(r.context["identity"]["nip"], "123456")


class HujanReportTest(TestCase):
    def setUp(self):
        self.client.force_login(get_user_model().objects.create_user("t", password="x"))

    def test_hujan_report(self):
        from hujan.models import Hujan
        Hujan.objects.create(tanggal=datetime.date(2026, 1, 1), hilman=4.4, obs=4.3, kategori="sangat ringan")
        Hujan.objects.create(tanggal=datetime.date(2026, 1, 2), hilman=9999, obs=0.4, kategori="Hujan Sangat  Ringan",
                             keterangan="Timer macet")
        Hujan.objects.create(tanggal=datetime.date(2026, 1, 5), hilman=86.3, obs=97.0, kategori="Hukan Lebat")
        Hujan.objects.create(tanggal=datetime.date(2026, 1, 6), hilman=3.9, obs=2.0, kategori="NIHIIL")
        r = self.client.get(report_url("hujan", start=(1, 2026), end=(1, 2026)))
        self.assertEqual(r.status_code, 200)
        c = r.context
        self.assertEqual(c["n_records"], 4)
        self.assertEqual(len(c["missing"]), 27)
        self.assertAlmostEqual(c["total"], 103.7)
        self.assertEqual(c["peak"]["obs"], 97.0)
        self.assertEqual([d["date"].day for d in c["hellmann_off"]], [2])
        self.assertEqual(dict(c["categories"]), {"Nihil": 1, "Sangat Ringan": 2, "Lebat": 1})
        self.assertEqual([d["date"].day for d in c["big_diff"]], [6])
        self.assertContains(r, "Tersedianya data curah hujan yang berkualitas")
        self.assertContains(r, "tidak beroperasi")

    def test_hujan_in_report_type_choices(self):
        self.assertContains(self.client.get("/reporting/"), "Hujan (Curah Hujan Harian)")


class HujanCreditTest(TestCase):
    """Hujan 07:00 WIT dikreditkan ke semua petugas dinas malam (MT/PSM) dari hari sebelumnya."""

    def setUp(self):
        from jadwal.models import JadwalHarian, Pegawai, PolaDinas
        self.MT = PolaDinas.objects.create(kode="MT", nama="Malam Tengah Malam", jam_mulai=datetime.time(19, 30),
                                           jam_selesai=datetime.time(8, 0), durasi=12.5)
        self.PSM = PolaDinas.objects.create(kode="PSM", nama="Pagi Siang Malam", jam_mulai=datetime.time(7, 30),
                                            jam_selesai=datetime.time(8, 0), durasi=24.5)
        self.PS = PolaDinas.objects.create(kode="PS", nama="Pagi Siang", jam_mulai=datetime.time(7, 30),
                                           jam_selesai=datetime.time(20, 0), durasi=12.5)
        U = get_user_model()
        self.juan_user = U.objects.create_user("juan", password="x")
        self.lidya_user = U.objects.create_user("lidya", password="x")
        self.juan = Pegawai.objects.create(user=self.juan_user, nama="GREGORIUS JUAN P. P. HALAWA, S.Tr.Geof", nip="1", pangkat="-", jabatan="-")
        self.lidya = Pegawai.objects.create(user=self.lidya_user, nama="LIDYA N. HUTAPEA, S.Si", nip="2", pangkat="-", jabatan="-")
        self.richard = Pegawai.objects.create(nama="RICHARD MILHOUS PANGKALI, S.Tr.Geof", nip="3", pangkat="-", jabatan="-")
        J = JadwalHarian.objects.create
        J(pegawai=self.juan, tanggal=datetime.date(2026, 1, 4), pola=self.MT)     # -> pembacaan 5 Jan
        J(pegawai=self.lidya, tanggal=datetime.date(2026, 1, 4), pola=self.MT)
        J(pegawai=self.richard, tanggal=datetime.date(2026, 1, 5), pola=self.PS)  # dinas pagi 5 Jan: tidak
        J(pegawai=self.juan, tanggal=datetime.date(2026, 1, 9), pola=self.PSM)    # -> pembacaan 10 Jan
        J(pegawai=self.lidya, tanggal=datetime.date(2026, 1, 9), pola=self.MT, keterangan_lain="CUTI")
        J(pegawai=self.lidya, tanggal=datetime.date(2026, 1, 19), pola=self.MT)   # -> 20 Jan, tercatat Richard
        from hujan.models import Hujan
        Hujan.objects.create(tanggal=datetime.date(2026, 1, 5), hilman=95, obs=97.0, kategori="sangat lebat", petugas="Juan")
        Hujan.objects.create(tanggal=datetime.date(2026, 1, 10), hilman=3, obs=3.2, kategori="sangat ringan", petugas="juan")
        Hujan.objects.create(tanggal=datetime.date(2026, 1, 20), hilman=0, obs=0, kategori="nihil", petugas="richard")

    def test_staff_on_duty_and_name_matching(self):
        from jadwal.duty import all_pegawai, match_pegawai, staff_on_duty_at
        t5, t10 = datetime.datetime(2026, 1, 5, 7), datetime.datetime(2026, 1, 10, 7)
        on = staff_on_duty_at([t5, t10])
        self.assertEqual(set(on[t5]), {self.juan, self.lidya})
        self.assertEqual(on[t10], [self.juan])  # PSM 24,5 jam; Lidya cuti
        peg = all_pegawai()
        self.assertEqual(match_pegawai("Juan", peg), self.juan)
        self.assertEqual(match_pegawai("RICHARD MILHOUS PANGKALI, S.Tr.Geof", peg), self.richard)
        self.assertIsNone(match_pegawai("staff ops", peg))

    def test_report_credits_both_night_officers_and_swaps(self):
        self.client.force_login(self.lidya_user)
        r = self.client.get(report_url("hujan", start=(1, 2026), end=(1, 2026)))
        c = r.context
        self.assertTrue(c["per_officer"])
        days = [d["date"].day for m in c["months"] for d in m["days"]]
        self.assertEqual(days, [5, 20])          # 5 Jan walau tercatat Juan; 20 Jan tukar dinas tercatat Richard
        jan5 = c["months"][0]["days"][0]
        self.assertEqual(set(jan5["officers"]), {self.juan.nama, self.lidya.nama})
        jan20 = c["months"][0]["days"][1]
        self.assertEqual(set(jan20["officers"]), {self.lidya.nama, self.richard.nama})
        self.assertEqual(c["n_records"], 2)
        self.assertContains(r, "LIDYA N. HUTAPEA")

        self.client.force_login(self.juan_user)
        c = self.client.get(report_url("hujan", start=(1, 2026), end=(1, 2026))).context
        self.assertEqual([d["date"].day for m in c["months"] for d in m["days"]], [5, 10])


class GempabumiReportTest(TestCase):
    @mock.patch("reporting.osm._tile", side_effect=fake_tile)
    def test_per_officer_events_map_and_list(self, _tile):
        from jadwal.models import Pegawai
        from qc_review.models import Event, QCRun
        user = get_user_model().objects.create_user("jambari", password="x")
        me = Pegawai.objects.create(user=user, nama="JAMBARI, S.Tr, M.Si", nip="1", pangkat="-", jabatan="-")
        other = Pegawai.objects.create(nama="ANAS", nip="2", pangkat="-", jabatan="-")
        def ev(pid, when, mag, depth, who):
            e = Event.objects.create(public_id=pid, origin_time=when, latitude=-3.0, longitude=139.5,
                                     depth_km=depth, magnitude=mag, region="Jayapura")
            QCRun.objects.create(event=e, run_number=1)
            e.petugas.set(who)
            return e
        ev("a", datetime.datetime(2026, 1, 5, 2, tzinfo=UTC), 4.4, 10, [me, other])
        ev("b", datetime.datetime(2026, 2, 7, 2, tzinfo=UTC), 5.4, 133, [me])
        ev("c", datetime.datetime(2026, 2, 8, 2, tzinfo=UTC), 3.0, 20, [other])
        ev("d", datetime.datetime(2026, 4, 1, 2, tzinfo=UTC), 3.1, 20, [me])        # di luar periode
        Event.objects.create(public_id="norun", origin_time=datetime.datetime(2026, 1, 9, tzinfo=UTC),
                             latitude=-3, longitude=139, depth_km=10, magnitude=3)  # belum di-commit
        self.client.force_login(user)
        r = self.client.get(report_url("gempabumi"))
        self.assertEqual(r.status_code, 200)
        c = r.context
        self.assertEqual([e["id"] for e in c["events"]], ["a", "b"])
        self.assertEqual(c["n_m5"], 1)
        self.assertEqual([m["n"] for m in c["monthly"]], [1, 1, 0])
        self.assertIsNotNone(c["map"])
        self.assertContains(r, "Tersedianya Informasi Gempabumi dan Tsunami yang dianalisis dan disebarkan")
        self.assertNotContains(r, "Delta")
        self.assertNotContains(r, "ON TIME")
        self.assertEqual(c["partners"], [("ANAS", 1)])
        # 5 Jan 02:00 UTC = 11:00 WIT; kedua kolom waktu tidak boleh sama
        self.assertContains(r, "<td>02:00:00</td><td>11:00:00</td>", html=False)

        anon = get_user_model().objects.create_user("tamu", password="x")
        self.client.force_login(anon)
        self.assertEqual(self.client.get(report_url("gempabumi")).context["n"], 3)


class GempabumiCepatReportTest(TestCase):
    @mock.patch("reporting.osm._tile", side_effect=fake_tile)
    def test_delta_and_percentages(self, _tile):
        from jadwal.models import Pegawai
        from qc_review.models import Event, QCRun
        user = get_user_model().objects.create_user("jambari", password="x")
        me = Pegawai.objects.create(user=user, nama="JAMBARI, S.Tr, M.Si", nip="1", pangkat="-", jabatan="-")
        def ev(pid, when, mag, delay_s, who=(me,)):
            e = Event.objects.create(public_id=pid, origin_time=when, latitude=-3.0, longitude=139.5,
                                     depth_km=10, magnitude=mag, region="Jayapura")
            Event.objects.filter(pk=e.pk).update(created_at=when + datetime.timedelta(seconds=delay_s))
            QCRun.objects.create(event=e, run_number=1)
            e.petugas.set(who)
        ev("old", datetime.datetime(2025, 12, 31, 2, tzinfo=UTC), 3.0, 900000)       # sebelum 2026: tanpa delta
        ev("fast", datetime.datetime(2026, 6, 5, 2, tzinfo=UTC), 3.5, 161)           # 02:41 tepat waktu
        ev("edge", datetime.datetime(2026, 6, 6, 2, tzinfo=UTC), 3.5, 180)           # tepat 3 menit = terlambat
        ev("slow", datetime.datetime(2026, 7, 6, 2, tzinfo=UTC), 4.9, 3725)          # 1:02:05 terlambat
        ev("big", datetime.datetime(2026, 6, 7, 2, tzinfo=UTC), 5.0, 60)             # M ≥ 5: tidak termasuk
        self.client.force_login(user)
        r = self.client.get(report_url("gempabumi_cepat", start=(12, 2025), end=(7, 2026)))
        self.assertEqual(r.status_code, 200)
        c = r.context
        self.assertEqual([e["id"] for e in c["events"]], ["old", "fast", "edge", "slow"])
        self.assertEqual((c["n_timed"], c["n_on"], c["n_late"]), (3, 1, 2))
        self.assertAlmostEqual(c["pct_on"], 100 / 3)
        self.assertEqual([m["timed"] for m in c["monthly"]], [0, 0, 0, 0, 0, 0, 2, 1])
        self.assertEqual([e["delta_text"] for e in c["events"]], ["–", "02:41", "03:00", "1:02:05"])
        self.assertIsNotNone(c["chart"])
        self.assertContains(r, "kurang dari 3 menit")
        self.assertContains(r, "33,3")
        self.assertContains(r, "<td>02:00:00</td><td>11:00:00</td>", html=False)

    @mock.patch("reporting.osm._tile", side_effect=fake_tile)
    def test_station_health_on_map_and_appendix(self, _tile):
        from qc_review.models import Event, QCRun, StationDailyQC, StationDiagnosis
        from stations.models import Station
        e = Event.objects.create(public_id="x", origin_time=datetime.datetime(2026, 6, 5, 2, tzinfo=UTC),
                                 latitude=-3.0, longitude=139.5, depth_km=10, magnitude=3.0)
        QCRun.objects.create(event=e, run_number=1)
        day = datetime.date(2026, 10, 17)
        for sta, chans in {"JAY": {"SHZ": "green", "SHN": "green"}, "ARMI": {"SHZ": "red", "SHN": "green"},
                           "AMPM": {"SHZ": "red", "SHN": "red"}}.items():
            for cha, sev in chans.items():
                StationDailyQC.objects.create(network="IA", station=sta, channel=cha, day=day, status=sev)
        StationDiagnosis.objects.create(network="IA", station="ARMI", channel="SHZ", day=day, code="COMPONENT_DEAD",
                                        severity="red", message_id="SHZ mati.", message_en="SHZ dead.")
        for cha in ("SHZ", "SHN"):
            StationDiagnosis.objects.create(network="IA", station="AMPM", channel=cha, day=day, code="NO_DATA",
                                            severity="red", message_id="Tidak ada data.", message_en="No data.")
        for code, lat in (("JAY", -2.51), ("ARMI", -2.9), ("AMPM", -3.3), ("DYPI", -4.85)):
            Station.objects.create(network="IA", code=code, latitude=lat, longitude=140.0,
                                   station_type=Station.TYPE_SEISMIC)
        self.client.force_login(get_user_model().objects.create_user("tamu", password="x"))
        r = self.client.get(report_url("gempabumi_cepat", start=(6, 2026), end=(6, 2026)))
        h = r.context["health"]
        self.assertEqual({s["station"]: s["state"] for s in h["stations"]},
                         {"JAY": "baik", "ARMI": "gangguan", "AMPM": "off", "DYPI": "off"})
        self.assertEqual([k for _lab, _c, k in h["counts"]], [1, 1, 2])
        self.assertIsNotNone(r.context["map"])
        self.assertContains(r, "Hasil Station Health Monitoring")
        self.assertContains(r, "V. Kendala")
        self.assertContains(r, "SHZ mati.")


class DeltaRuleTest(TestCase):
    def test_no_on_time_default_before_2026(self):
        from qc_review.utils import delta_on_time_info
        old = datetime.datetime(2025, 12, 1, tzinfo=UTC)
        self.assertEqual(delta_on_time_info(old, old + datetime.timedelta(days=90)), (None, "–"))
        t = datetime.datetime(2026, 2, 1, tzinfo=UTC)
        self.assertEqual(delta_on_time_info(t, t + datetime.timedelta(minutes=7)), (7.0, "LATE"))
        self.assertEqual(delta_on_time_info(t, t + datetime.timedelta(minutes=3)), (3.0, "LATE"))
        self.assertEqual(delta_on_time_info(t, t + datetime.timedelta(seconds=161)), (2.7, "ON TIME"))


class TandaWaktuReportTest(TestCase):
    @mock.patch("reporting.osm._tile", side_effect=fake_tile)
    def test_times_grouped_by_wit_date(self, _tile):
        from almanac.models import SunMoonEvent
        # Satu baris per tanggal WIT, waktu UTC: terbit 20:40 UTC tgl 14 = 05:40 WIT tgl 15.
        for day in range(14, 18):
            SunMoonEvent.objects.create(
                date=datetime.date(2026, 1, day), city="Jayapura",
                sun_rise=datetime.datetime(2026, 1, day - 1, 20, 40, tzinfo=UTC),
                sun_set=datetime.datetime(2026, 1, day, 8, 52, tzinfo=UTC),
                moon_rise=datetime.datetime(2026, 1, day, 3, 5, tzinfo=UTC) if day != 16 else None,
                moon_set=datetime.datetime(2026, 1, day, 14, 30, tzinfo=UTC), moon_phase=0.5)
        self.client.force_login(get_user_model().objects.create_user("tamu", password="x"))
        r = self.client.get(report_url("tanda_waktu", start=(1, 2026), end=(1, 2026)))
        self.assertEqual(r.status_code, 200)
        c = r.context
        jan = {row["date"].day: row for row in c["home_months"][0]["rows"]}
        self.assertEqual((jan[15]["rise"], jan[15]["set"], jan[15]["len"]), ("05:40", "17:52", "12 j 12 m"))
        self.assertEqual((jan[15]["mrise"], jan[15]["mset"]), ("12:05", "23:30"))
        self.assertEqual(jan[16]["mrise"], "–")
        self.assertIn("Purnama", jan[15]["phase"])
        self.assertEqual(jan[13]["rise"], "–")
        jay = next(x for x in c["cities"] if x["name"] == "Jayapura")
        self.assertEqual(jay["n"], 4)
        self.assertEqual(c["n_days"], 31)
        self.assertIsNotNone(c["map"])
        self.assertIsNotNone(c["chart"])
        self.assertContains(r, "Terlaksananya layanan informasi tanda waktu di Wilayah Papua dan Sekitarnya")
